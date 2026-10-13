"""Word (.docx) hujjatni formatini saqlab tarjima qilish.

Har bir abzats va jadval katagi matni butunlay tarjima qilinib, birinchi matnli run'ga yoziladi — shunda
abzats uslubi, shrift, jadval tuzilmasi, rasmlar va sahifa sozlamalari o'z joyida qoladi. Abzats ichidagi
aralash formatlash (bir so'z qalin, qolgani oddiy) birinchi run uslubiga tenglashadi."""
from __future__ import annotations

import io
import re
from typing import Awaitable, Callable

MAX_SEGMENTS = 1500
BATCH_ITEMS = 40
BATCH_CHARS = 6000

Translator = Callable[[list[str]], Awaitable[list[str]]]
Progress = Callable[[int, int], Awaitable[None]]


class TranslateError(Exception):
    pass


def _walk(container):
    """Konteyner (hujjat, katak, sarlavha/izoh) ichidagi abzatslar — jadvallar ichiga ham kiradi."""
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        seen = set()  # birlashtirilgan katak (eniga ham, bo'yiga ham) bir marta — har qatorda qayta kelmasin
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield from _walk(cell)


def _translatable(paragraph) -> bool:
    text = paragraph.text
    if not text.strip() or not re.search(r"[^\W\d_]", text):  # harf yo'q — raqam, sana, belgi
        return False
    # Maydonlar (sahifa raqami, mundarija) — Word o'zi hisoblaydi, tegmaymiz
    return not paragraph._p.xpath(".//w:fldChar | .//w:instrText | .//w:fldSimple")


def collect(document) -> list:
    paragraphs = list(_walk(document))
    for section in document.sections:
        for part in (
            section.header, section.footer, section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if part.is_linked_to_previous:
                continue
            paragraphs.extend(_walk(part))
    return [p for p in paragraphs if _translatable(p)]


def _runs(paragraph) -> list:
    from docx.text.hyperlink import Hyperlink
    from docx.text.run import Run

    out = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Run):
            out.append(item)
        elif isinstance(item, Hyperlink):
            out.extend(item.runs)
    return out


def set_text(paragraph, text: str) -> None:
    runs = [r for r in _runs(paragraph) if r.text]
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        if run._r.xpath(".//w:drawing | .//w:pict"):  # rasmli run — tegmaymiz
            continue
        run.text = ""


def batches(texts: list[str]) -> list[list[str]]:
    groups: list[list[str]] = []
    current: list[str] = []
    size = 0
    for text in texts:
        if current and (len(current) >= BATCH_ITEMS or size + len(text) > BATCH_CHARS):
            groups.append(current)
            current, size = [], 0
        current.append(text)
        size += len(text)
    if current:
        groups.append(current)
    return groups


async def translate_docx(data: bytes, translate: Translator, progress: Progress | None = None) -> tuple[bytes, int]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise TranslateError("Word faylni o'qib bo'lmadi (faqat .docx).") from exc
    paragraphs = collect(document)
    if not paragraphs:
        raise TranslateError("Hujjatda tarjima qilinadigan matn topilmadi.")
    if len(paragraphs) > MAX_SEGMENTS:
        raise TranslateError(f"Hujjat juda katta ({len(paragraphs)} abzats); {MAX_SEGMENTS} tagacha bo'lsin.")
    texts = [p.text for p in paragraphs]
    translated: list[str] = []
    for group in batches(texts):
        result = await translate(group)
        if len(result) != len(group):
            raise TranslateError("Tarjima natijasi bo'laklar soniga mos kelmadi — qayta urinib ko'ring.")
        translated.extend(str(item) if item is not None else "" for item in result)
        if progress is not None:
            await progress(len(translated), len(texts))
    for paragraph, text in zip(paragraphs, translated):
        set_text(paragraph, text or paragraph.text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), len(paragraphs)
