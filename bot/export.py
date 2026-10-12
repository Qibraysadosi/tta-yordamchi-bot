"""Javobni Word (.docx) hujjatiga aylantirish — Markdown belgilari sarlavha, ro'yxat va qalin matnga o'tadi."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt

_HEADER = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^\s*[-*+•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_SEP_ROW = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`|\*[^*\n]+\*|_[^_\n]+_)")


def _add_runs(paragraph, text: str) -> None:
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if (piece.startswith("**") and piece.endswith("**")) or (piece.startswith("__") and piece.endswith("__")):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif len(piece) > 2 and ((piece.startswith("*") and piece.endswith("*")) or (piece.startswith("_") and piece.endswith("_"))):
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def _cells(row: str) -> list[str]:
    return [c.strip() for c in row.strip().strip("|").split("|")]


def answer_to_docx(text: str, title: str | None = None) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if title:
        doc.add_heading(title, level=1)

    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    in_code = False
    code_lines: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                code_lines = []
            in_code = not in_code
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if _TABLE_ROW.match(line):
            rows = []
            while i < len(lines) and _TABLE_ROW.match(lines[i]):
                if not _SEP_ROW.match(lines[i]):
                    rows.append(_cells(lines[i]))
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(width):
                        cell = table.cell(r, c)
                        cell.text = ""
                        _add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        header = _HEADER.match(line)
        if header:
            doc.add_heading(header.group(2).replace("**", ""), level=min(len(header.group(1)) + 1, 4))
        elif (m := _BULLET.match(line)):
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif (m := _NUMBERED.match(line)):
            _add_runs(doc.add_paragraph(style="List Number"), m.group(2))
        elif line.strip():
            _add_runs(doc.add_paragraph(), line.strip())
        i += 1
    if in_code and code_lines:
        doc.add_paragraph("\n".join(code_lines))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
