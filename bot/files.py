"""Fayl yaratuvchilar: Word, Excel, PowerPoint, PDF, grafik — Claude vositalari uchun (bot/tools.py)."""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

from .export import answer_to_docx

FONTS = Path(__file__).resolve().parent / "assets" / "fonts"
PALETTE = ["#4F46E5", "#06B6D4", "#F59E0B", "#10B981", "#EF4444", "#8B5CF6", "#64748B", "#EC4899"]
MAX_ROWS = 5000


class FileError(Exception):
    """Vosita xatosi — Claude'ga tool_result sifatida qaytariladi, u foydalanuvchiga tushuntiradi."""


def safe_name(name: str, extension: str) -> str:
    base = re.sub(r"[^\w.\- ]+", "_", (name or "fayl").strip(), flags=re.UNICODE).strip(" ._") or "fayl"
    base = base[:80]
    if not base.lower().endswith(f".{extension}"):
        base = f"{base}.{extension}"
    return base


# ------------------------------------------------------------------ Word / matn
def build_docx(content: str, title: str | None = None) -> bytes:
    return answer_to_docx(content, title)


def build_text(content: str) -> bytes:
    return content.encode("utf-8")


# ------------------------------------------------------------------ Excel
def _cell_value(value: Any) -> Any:
    if isinstance(value, str):
        text = value.strip()
        if text and re.fullmatch(r"-?\d+", text):
            return int(text)
        if text and re.fullmatch(r"-?\d+[.,]\d+", text):
            return float(text.replace(",", "."))
        return value
    return value


def build_xlsx(sheets: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not sheets:
        raise FileError("Kamida bitta varaq (sheet) kerak.")
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for index, sheet in enumerate(sheets, 1):
        name = re.sub(r"[\[\]:*?/\\]", " ", str(sheet.get("name") or f"Varaq{index}")).strip()[:28] or f"Varaq{index}"
        if name.lower() in used:  # bir xil nomli varaqlar bo'lmasin
            name = f"{name}_{index}"
        used.add(name.lower())
        ws = wb.create_sheet(title=name)
        columns = list(sheet.get("columns") or [])
        rows = list(sheet.get("rows") or [])
        if len(rows) > MAX_ROWS:
            raise FileError(f"Juda ko'p qator ({len(rows)}); {MAX_ROWS} tagacha bo'lsin.")
        if columns:
            ws.append([str(c) for c in columns])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4F46E5")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
        for row in rows:
            if not isinstance(row, (list, tuple)):
                raise FileError("Har bir qator ro'yxat bo'lishi kerak: [[...], [...]].")
            ws.append([_cell_value(v) for v in row])
        widths: dict[int, int] = {}
        for row in ws.iter_rows(values_only=True):
            for col, value in enumerate(row, 1):
                length = len(str(value)) if value is not None else 0
                widths[col] = max(widths.get(col, 0), min(length, 60))
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = max(8, width + 2)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_csv(sheets: list[dict]) -> bytes:
    if not sheets:
        raise FileError("Kamida bitta varaq kerak.")
    sheet = sheets[0]
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    if sheet.get("columns"):
        writer.writerow(sheet["columns"])
    for row in sheet.get("rows") or []:
        writer.writerow(row)
    return ("﻿" + buffer.getvalue()).encode("utf-8")


# ------------------------------------------------------------------ PowerPoint
def build_pptx(slides: list[dict], title: str | None = None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    if not slides:
        raise FileError("Kamida bitta slayd kerak.")
    if len(slides) > 60:
        raise FileError("Slaydlar 60 tadan oshmasin.")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    dark, accent, white, grey = RGBColor(0x1E, 0x1B, 0x4B), RGBColor(0x4F, 0x46, 0xE5), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x47, 0x55, 0x69)
    blank = prs.slide_layouts[6]

    def add_box(slide, left, top, width, height, text, size, bold=False, color=white, align=None):
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.text = text
        paragraph.font.size, paragraph.font.bold, paragraph.font.color.rgb = Pt(size), bold, color
        if align is not None:
            paragraph.alignment = align
        return frame

    if title:
        slide = prs.slides.add_slide(blank)
        bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid(); bg.fill.fore_color.rgb = dark; bg.line.fill.background()
        bar = slide.shapes.add_shape(1, Inches(0.8), Inches(3.55), Inches(1.2), Inches(0.08))
        bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
        add_box(slide, Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.5), title, 44, bold=True)
        subtitle = slides[0].get("subtitle") if isinstance(slides[0], dict) else None
        if subtitle:
            add_box(slide, Inches(0.8), Inches(3.9), Inches(11.5), Inches(1.0), str(subtitle), 22, color=RGBColor(0xC7, 0xD2, 0xFE))

    for item in slides:
        if not isinstance(item, dict):
            raise FileError("Har slayd {title, bullets} ko'rinishida bo'lsin.")
        slide = prs.slides.add_slide(blank)
        header = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.25))
        header.fill.solid(); header.fill.fore_color.rgb = dark; header.line.fill.background()
        add_box(slide, Inches(0.6), Inches(0.28), Inches(12), Inches(0.9), str(item.get("title") or ""), 30, bold=True)
        bullets = [str(b) for b in (item.get("bullets") or [])]
        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(12), Inches(5.3)).text_frame
        body.word_wrap = True
        for i, bullet in enumerate(bullets[:12]):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = f"•  {bullet}"
            paragraph.font.size, paragraph.font.color.rgb = Pt(20 if len(bullets) <= 6 else 16), grey
            paragraph.space_after = Pt(10)
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(item["notes"])
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ PDF
_MD_HEADER = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_BULLET = re.compile(r"^\s*[-*+•]\s+(.*)$")
_MD_NUMBER = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")
_MD_TABLE = re.compile(r"^\s*\|.*\|\s*$")
_MD_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_MD_INLINE = re.compile(r"\*\*(.+?)\*\*|`([^`]+)`")


def _plain(text: str) -> str:
    return _MD_INLINE.sub(lambda m: m.group(1) or m.group(2), text).replace("__", "")


def build_pdf(content: str, title: str | None = None) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_font("Sans", "", str(FONTS / "DejaVuSans.ttf"))
    pdf.add_font("Sans", "B", str(FONTS / "DejaVuSans-Bold.ttf"))
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def paragraph(text: str, size: float = 11, bold: bool = False, indent: float = 0, after: float = 2.5) -> None:
        pdf.set_font("Sans", "B" if bold else "", size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(width - indent, size * 0.55, _plain(text))
        pdf.ln(after)

    if title:
        paragraph(title, 18, bold=True, after=5)
    lines = content.replace("\r\n", "\n").split("\n")
    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            pdf.set_font("Sans", "", 9)
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(width - 4, 5, line or " ")
            i += 1
            continue
        if _MD_TABLE.match(line):
            rows = []
            while i < len(lines) and _MD_TABLE.match(lines[i]):
                if not _MD_SEP.match(lines[i]):
                    rows.append([_plain(c.strip()) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                pdf.set_font("Sans", "", 9.5)
                with pdf.table(width=width, col_widths=[width / cols] * cols, text_align="LEFT") as table:
                    for r_index, row in enumerate(rows):
                        tr = table.row()
                        for c in range(cols):
                            tr.cell(row[c] if c < len(row) else "")
                pdf.ln(3)
            continue
        if (m := _MD_HEADER.match(line)):
            level = len(m.group(1))
            paragraph(m.group(2), {1: 16, 2: 14, 3: 12.5}.get(level, 11.5), bold=True, after=2)
        elif (m := _MD_BULLET.match(line)):
            paragraph("•  " + m.group(1), indent=4, after=1)
        elif (m := _MD_NUMBER.match(line)):
            paragraph(f"{m.group(1)}.  {m.group(2)}", indent=4, after=1)
        elif line.strip():
            paragraph(line.strip())
        else:
            pdf.ln(2)
        i += 1
    return bytes(pdf.output())


# ------------------------------------------------------------------ Markdown'dan jadval / slaydlar
def tables_from_markdown(content: str) -> list[dict]:
    """Matndagi Markdown jadvallarini varaqlarga aylantiradi (sheets berilmaganda zaxira yo'l)."""
    sheets: list[dict] = []
    lines = (content or "").replace("\r\n", "\n").split("\n")
    i = 0
    heading = None
    while i < len(lines):
        line = lines[i]
        if (m := _MD_HEADER.match(line)):
            heading = _plain(m.group(2))
        if _MD_TABLE.match(line):
            rows: list[list[str]] = []
            while i < len(lines) and _MD_TABLE.match(lines[i]):
                if not _MD_SEP.match(lines[i]):
                    rows.append([_plain(c.strip()) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            if rows:
                sheets.append({"name": heading or f"Varaq{len(sheets) + 1}", "columns": rows[0], "rows": rows[1:]})
                heading = None
            continue
        i += 1
    return sheets


def slides_from_markdown(content: str) -> list[dict]:
    """'# Sarlavha' + '- band' ko'rinishidagi matnni slaydlarga bo'ladi (slides berilmaganda zaxira yo'l)."""
    slides: list[dict] = []
    current: dict | None = None
    for line in (content or "").replace("\r\n", "\n").split("\n"):
        if (m := _MD_HEADER.match(line)):
            current = {"title": _plain(m.group(2)), "bullets": []}
            slides.append(current)
        elif (m := _MD_BULLET.match(line)) or (m := _MD_NUMBER.match(line)):
            if current is None:
                current = {"title": "", "bullets": []}
                slides.append(current)
            current["bullets"].append(_plain(m.group(m.lastindex)))
        elif line.strip() and current is not None:
            current["bullets"].append(_plain(line.strip()))
    return slides


# ------------------------------------------------------------------ grafik
def build_chart(spec: dict) -> bytes:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    kind = str(spec.get("type") or "bar").lower()
    labels = [str(x) for x in (spec.get("labels") or [])]
    series = spec.get("series") or []
    if not labels or not series:
        raise FileError("Grafik uchun labels va series kerak.")
    for s in series:
        values = s.get("values") or []
        if len(values) != len(labels):
            raise FileError("Har seriyadagi qiymatlar soni labels soniga teng bo'lsin.")
        try:
            s["values"] = [float(v) for v in values]
        except (TypeError, ValueError) as exc:
            raise FileError("Qiymatlar son bo'lishi kerak.") from exc
    if len(labels) > 60:
        raise FileError("60 tadan ko'p nuqta — ko'rinmaydi; qisqartiring.")

    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=150)
    fig.patch.set_facecolor("white")
    if kind == "pie":
        values = series[0]["values"]
        ax.pie(values, labels=labels, autopct="%1.0f%%", colors=PALETTE[: len(values)], startangle=90,
               wedgeprops={"linewidth": 1.5, "edgecolor": "white"}, textprops={"fontsize": 10})
        ax.axis("equal")
    else:
        x = range(len(labels))
        if kind == "line":
            for i, s in enumerate(series):
                ax.plot(list(x), s["values"], marker="o", linewidth=2.2, color=PALETTE[i % len(PALETTE)], label=str(s.get("name") or f"{i + 1}"))
        elif kind == "barh":
            for i, s in enumerate(series):
                ax.barh([p + i * 0.8 / len(series) for p in x], s["values"], height=0.8 / len(series), color=PALETTE[i % len(PALETTE)], label=str(s.get("name") or f"{i + 1}"))
            ax.set_yticks([p + 0.4 - 0.4 / len(series) for p in x])
            ax.set_yticklabels(labels)
        else:
            n = len(series)
            w = 0.8 / n
            for i, s in enumerate(series):
                bars = ax.bar([p + i * w - 0.4 + w / 2 for p in x], s["values"], width=w, color=PALETTE[i % len(PALETTE)], label=str(s.get("name") or f"{i + 1}"))
                if len(labels) <= 15 and n <= 2:
                    ax.bar_label(bars, fmt="%g", fontsize=8, padding=2)
        if kind != "barh":
            ax.set_xticks(list(x))
            ax.set_xticklabels(labels, rotation=30 if max(len(l) for l in labels) > 8 else 0, ha="right" if max(len(l) for l in labels) > 8 else "center")
        ax.grid(axis="x" if kind == "barh" else "y", alpha=0.25)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        if len(series) > 1 or series[0].get("name"):
            ax.legend(frameon=False)
    if spec.get("title"):
        ax.set_title(str(spec["title"]), fontsize=14, fontweight="bold", pad=14)
    if spec.get("x_label"):
        ax.set_xlabel(str(spec["x_label"]))
    if spec.get("y_label"):
        ax.set_ylabel(str(spec["y_label"]))
    fig.tight_layout()
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png")
    plt.close(fig)
    return buffer.getvalue()
