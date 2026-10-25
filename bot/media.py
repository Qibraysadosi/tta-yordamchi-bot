"""Telegram'dan kelgan rasm va fayllarni Claude tushunadigan bloklarga aylantirish."""
from __future__ import annotations

import base64
import codecs
import io
import os

from PIL import Image, ImageOps

MAX_TEXT_CHARS = 400_000  # ≈ 100 ming token

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
IMAGE_MIMES = {"image/jpeg", "image/png", "image/webp", "image/gif"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEXT_MIMES = {"application/json", "application/xml", "application/x-yaml", "application/x-sh", "application/sql"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".css", ".js", ".jsx", ".ts", ".tsx",
    ".py", ".java", ".kt", ".swift", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php", ".rb", ".sh",
    ".sql", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log", ".srt", ".vue", ".svelte", ".astro",
}


class MediaError(Exception):
    """Foydalanuvchiga ko'rsatiladigan xato matni."""


def classify(filename: str, mime: str) -> str | None:
    name = (filename or "").lower()
    ext = os.path.splitext(name)[1]
    mime = (mime or "").lower()
    if mime in IMAGE_MIMES or ext in IMAGE_EXTENSIONS:
        return "image"
    if mime == "application/pdf" or ext == ".pdf":
        return "pdf"
    if mime == DOCX_MIME or ext == ".docx":
        return "docx"
    if mime == XLSX_MIME or ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if mime.startswith("text/") or mime in TEXT_MIMES or ext in TEXT_EXTENSIONS:
        return "text"
    return None


MAX_PIXELS = 60_000_000  # JPEG bo'lmagan rasm uchun (60 MP): kattasi xotirani to'ldirib yubormasin
PNG_LIMIT = 3_500_000  # shaffof rasm PNG'da bundan katta bo'lsa — oq fonda JPEG (API'ga va tarixga og'ir)


def image_block(data: bytes, max_edge: int) -> dict:
    try:
        img = Image.open(io.BytesIO(data))
        img.seek(0)  # GIF bo'lsa — birinchi kadr
        if img.format == "JPEG":
            img.draft("RGB", (max_edge, max_edge))  # katta foto kichraytirib o'qiladi (xotira va vaqt tejaladi)
        elif img.width * img.height > MAX_PIXELS:
            raise MediaError(f"Rasm juda katta ({img.width}×{img.height}) — kichikroq qilib yuboring.")
        img.load()
        img = ImageOps.exif_transpose(img) or img
    except MediaError:
        raise
    except Exception as exc:
        raise MediaError("Rasmni o'qib bo'lmadi — fayl buzilgan yoki formati noma'lum.") from exc
    if img.mode in ("I", "I;16", "I;16B", "I;16L"):  # 16 bitli kulrang — 8 bitga (aks holda oppoq bo'ladi)
        img = img.convert("I").point(lambda v: v * (1 / 256)).convert("L")
    has_alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
    img = img.convert("RGBA" if has_alpha else "RGB")
    if has_alpha and img.getchannel("A").getextrema()[0] == 255:  # alfa kanal bor, lekin hammasi shaffofsiz
        img, has_alpha = img.convert("RGB"), False
    if max(img.size) > max_edge:
        img.thumbnail((max_edge, max_edge), Image.Resampling.LANCZOS)
    data_out, media_type = b"", "image/jpeg"
    if has_alpha:
        buf = io.BytesIO()
        img.save(buf, format="PNG", optimize=True)
        if buf.tell() <= PNG_LIMIT:
            data_out, media_type = buf.getvalue(), "image/png"
        else:  # foto kabi shaffof rasm — PNG juda og'ir: oq fon ustida JPEG
            background = Image.new("RGB", img.size, (255, 255, 255))
            background.paste(img, mask=img.getchannel("A"))
            img = background
    if not data_out:
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=85, optimize=True)
        data_out = buf.getvalue()
    return {
        "type": "image",
        "source": {"type": "base64", "media_type": media_type, "data": base64.b64encode(data_out).decode("ascii")},
    }


def pdf_block(data: bytes, filename: str) -> dict:
    if b"%PDF" not in data[:1024]:
        raise MediaError(f"«{filename}» PDF faylga o'xshamaydi.")
    return {
        "type": "document",
        "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode("ascii")},
        "title": filename[:200],
    }


def text_block(text: str, filename: str) -> dict:
    text = text.strip()
    if not text:
        raise MediaError(f"«{filename}» ichida matn topilmadi.")
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + f"\n\n[… fayl juda uzun — dastlabki {MAX_TEXT_CHARS:,} belgi olindi]"
    return {
        "type": "document",
        "source": {"type": "text", "media_type": "text/plain", "data": text},
        "title": filename[:200],
    }


def decode_text(data: bytes) -> str:
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8):].decode("utf-8", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    for encoding in ("utf-8", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def docx_text(data: bytes) -> str:
    try:
        import docx
        from docx.table import Table

        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise MediaError("Word faylni o'qib bo'lmadi (faqat .docx formati qo'llab-quvvatlanadi).") from exc
    lines: list[str] = []
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                cells, previous = [], None
                for cell in row.cells:
                    if cell._tc is previous:  # birlashtirilgan katak takrorlanmasin
                        continue
                    previous = cell._tc
                    cells.append(cell.text.strip().replace("\n", " "))
                lines.append(" | ".join(cells))
        else:
            lines.append(item.text)
    return "\n".join(lines)


def xlsx_text(data: bytes, max_rows: int = 3000) -> str:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise MediaError("Excel faylni o'qib bo'lmadi (faqat .xlsx formati qo'llab-quvvatlanadi).") from exc
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"### Varaq: {sheet.title}")
            count = 0
            for row in sheet.iter_rows(values_only=True):
                if not row or all(v is None or str(v).strip() == "" for v in row):
                    continue
                lines.append("\t".join("" if v is None else str(v) for v in row).rstrip("\t"))
                count += 1
                if count >= max_rows:
                    lines.append(f"[… varaqda {max_rows} qatordan ko'p, qolgani olinmadi]")
                    break
    finally:
        workbook.close()
    return "\n".join(lines)


def build_block(kind: str, data: bytes, filename: str, max_edge: int) -> dict:
    if kind == "image":
        return image_block(data, max_edge)
    if kind == "pdf":
        return pdf_block(data, filename)
    if kind == "docx":
        return text_block(docx_text(data), filename)
    if kind == "xlsx":
        return text_block(xlsx_text(data), filename)
    if kind == "text":
        return text_block(decode_text(data), filename)
    raise MediaError(f"«{filename}» turidagi faylni o'qiy olmayman.")
