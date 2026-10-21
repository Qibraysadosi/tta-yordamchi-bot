"""Haftalik va oylik Excel hisobot (v1.9): «Umumiy», «Kunlar», «Lidlar», «Vazifalar» varaqlari.

Ma'lumot bot/reports/archive.py da yig'iladi (ExcelReport — tayyor qiymatlar), bu yerda faqat fayl yasaladi:
og'ir ish, asyncio.to_thread da chaqiriladi. Rahbarlar guruhi nusxasida (masked=True) telefon raqamlari
yashiriladi — ism, tuman, izoh, vazifa matni kabi har qanday matndagilari ham. Formulalar yo'q: qiymatlar tayyor
(telefonda ham ko'rinadi); «=» bilan boshlangan matn formula bo'lib qolmaydi, Excel qabul qilmaydigan boshqaruv
belgilari olib tashlanadi."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from typing import Any

from .leads import mask_phones

FONT = "Arial"
BRAND = "4F46E5"
TOTAL_FILL = "EEF2FF"
GOOD_FILL = "E8F5E9"
WARN_FILL = "FFF4E5"
MUTED = "6B7280"
NOTE_LIMIT = 300
CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")  # Excel (openpyxl) qabul qilmaydigan belgilar
# Odatiy shakldan tashqari raqamlar ham (90 12 34 567, 0901234567, 90/123/45/67, +7 …): 9–15 ta raqam.
# Sana (21.10.2026) va summa (1 250 000) 9 raqamdan kam — tegilmaydi; «19:00» dagi soat qismiga yopishmaydi.
LOOSE_NUMBER_RE = re.compile(r"(?<![\w+])\+?(?:\d{9,15}|\d{1,4}(?:[\s\-./()]{1,2}\d{1,4}){2,6})(?![\d:])")


@dataclass
class OpLine:
    name: str
    username: str | None
    total: int
    counts: dict[str, int]
    positive: int
    listed: bool = True  # False — ro'yxatda yo'q, lekin qo'ng'iroq qilgan


@dataclass
class DayLine:
    day: date
    weekday: str
    total: int
    counts: dict[str, int]
    positive: int
    new_leads: int
    unmarked: int
    live: bool = False


@dataclass
class LeadLine:
    created: datetime | None  # mahalliy vaqt (tzinfo'siz — Excel vaqt zonasini bilmaydi)
    name: str
    phones: list[str]
    age: str
    district: str
    course: str
    status: str  # «👍 keladi» yoki «belgilanmagan»
    positive: bool
    people: str
    note: str
    reason: str
    callback: str
    link: str | None


@dataclass
class TaskLine:
    who: str
    text: str
    created: datetime | None
    due: datetime | None
    state: str
    closed: datetime | None
    no_count: int
    repeat: str


@dataclass
class ExcelReport:
    title: str
    period: str
    generated: str
    filename: str
    columns: list[tuple[str, str]]  # (emoji, ma'nosi) — status ustunlari
    positive_label: str
    summary: list[tuple[str, int, float | None]]  # (ko'rsatkich, son, ulush 0..1)
    operators: list[OpLine]
    courses: list[tuple[str, int, int]]
    reasons: list[tuple[str, int]]
    days: list[DayLine]
    leads: list[LeadLine]
    tasks: list[TaskLine]
    notes: list[str] = field(default_factory=list)


# ------------------------------------------------------------------ telefon
def _digits(phone: str) -> str:
    return re.sub(r"\D", "", phone or "")


def format_phone(phone: str) -> str:
    digits = _digits(phone)
    if len(digits) == 12 and digits.startswith("998"):
        return f"+998 {digits[3:5]} {digits[5:8]}-{digits[8:10]}-{digits[10:]}"
    return phone


def mask_phone(phone: str) -> str:
    """Rahbarlar nusxasi: operator kodi va oxirgi 2 raqam qoladi («+998 90 ***-**-67»)."""
    digits = _digits(phone)
    if len(digits) == 12 and digits.startswith("998"):
        return f"+998 {digits[3:5]} ***-**-{digits[10:]}"
    return "***"


def hide_numbers(text: str) -> str:
    """Matndagi telefon raqamlari (odatiy va noodatiy yozilganlari) → «[raqam]»."""
    def repl(match: re.Match) -> str:
        return "[raqam]" if 9 <= len(_digits(match.group(0))) <= 15 else match.group(0)
    return LOOSE_NUMBER_RE.sub(repl, mask_phones(text or ""))


def masked_report(report: "ExcelReport") -> "ExcelReport":
    """Rahbarlar nusxasi uchun: odamlar yozgan har bir matn maydonidan raqamlar olib tashlanadi."""
    hide = hide_numbers
    return replace(
        report,
        operators=[replace(o, name=hide(o.name), username=hide(o.username) if o.username else o.username)
                   for o in report.operators],
        courses=[(hide(course), count, positive) for course, count, positive in report.courses],
        reasons=[(hide(reason), count) for reason, count in report.reasons],
        leads=[replace(lead, name=hide(lead.name), age=hide(lead.age), district=hide(lead.district),
                       course=hide(lead.course), people=hide(lead.people), note=hide(lead.note),
                       reason=hide(lead.reason), callback=hide(lead.callback)) for lead in report.leads],
        tasks=[replace(task, who=hide(task.who), text=hide(task.text)) for task in report.tasks],
    )


# ------------------------------------------------------------------ yasash
class _Book:
    """openpyxl ustidan yupqa qatlam: bir xil shrift, sarlavha va jami qatorlari uslubi."""

    def __init__(self) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self.wb = Workbook()
        self.font = Font(name=FONT, size=10)
        self.bold = Font(name=FONT, size=10, bold=True)
        self.italic = Font(name=FONT, size=10, italic=True)
        self.head_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        self.title_font = Font(name=FONT, size=14, bold=True)
        self.section_font = Font(name=FONT, size=11, bold=True, color=BRAND)
        self.muted = Font(name=FONT, size=9, italic=True, color=MUTED)
        self.link_font = Font(name=FONT, size=10, color="2563EB", underline="single")
        self.head_fill = PatternFill("solid", fgColor=BRAND)
        self.total_fill = PatternFill("solid", fgColor=TOTAL_FILL)
        self.good_fill = PatternFill("solid", fgColor=GOOD_FILL)
        self.warn_fill = PatternFill("solid", fgColor=WARN_FILL)
        self.center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.top = Alignment(vertical="top", wrap_text=True)
        self.line = Border(bottom=Side(style="thin", color="E5E7EB"))

    def sheet(self, title: str, widths: list[float]):
        from openpyxl.utils import get_column_letter

        if self.wb.worksheets and self.wb.worksheets[0].title == "Sheet":
            ws = self.wb.worksheets[0]
            ws.title = title
        else:
            ws = self.wb.create_sheet(title)
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"  # chop etishda jadval bir sahifa kengligiga sig'sin
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        return ws

    def put(self, ws, row: int, col: int, value: Any, *, font=None, fill=None, fmt: str | None = None,
            align=None, border: bool = False):
        if isinstance(value, str):
            value = CONTROL_RE.sub("", value)
        cell = ws.cell(row=row, column=col, value=value)
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"  # lid formasidan kelgan «=HYPERLINK(…)» kabi matn formula bo'lmasin
        cell.font = font or self.font
        if fill is not None:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        if align is not None:
            cell.alignment = align
        if border:
            cell.border = self.line
        return cell

    def header(self, ws, row: int, titles: list[str]) -> None:
        for col, title in enumerate(titles, 1):
            self.put(ws, row, col, title, font=self.head_font, fill=self.head_fill, align=self.center)
        ws.row_dimensions[row].height = 40  # uzun status nomlari 3 qatorga sig'adi

    def values(self, ws, row: int, values: list[Any], *, font=None, fill=None, fmts: dict[int, str] | None = None,
               wrap: bool = False) -> None:
        for col, value in enumerate(values, 1):
            self.put(ws, row, col, value, font=font, fill=fill, fmt=(fmts or {}).get(col),
                     align=self.top if wrap else None, border=True)

    def save(self) -> bytes:
        self.wb.properties.creator = "TTA Yordamchi"
        buffer = io.BytesIO()
        self.wb.save(buffer)
        return buffer.getvalue()


def _share(part: int, whole: int) -> float | None:
    return part / whole if whole else None


def _status_titles(report: ExcelReport) -> list[str]:
    return [f"{emoji} {label}".strip() for emoji, label in report.columns]


def _summary_sheet(book: _Book, report: ExcelReport, masked: bool) -> None:
    cols = len(report.columns)
    ws = book.sheet("Umumiy", [30, 16] + [12] * (cols + 2))
    book.put(ws, 1, 1, report.title, font=book.title_font)
    book.put(ws, 2, 1, report.period, font=book.font)
    row = 3
    for note in [report.generated] + report.notes + (
        ["🔒 Rahbarlar nusxasi: telefon raqamlari yashirilgan"] if masked else ["🔓 To'liq nusxa (telefon raqamlari bilan)"]
    ):
        book.put(ws, row, 1, note, font=book.muted)
        row += 1

    row += 1
    book.put(ws, row, 1, "Asosiy ko'rsatkichlar", font=book.section_font)
    row += 1
    for label, number, share in report.summary:
        book.put(ws, row, 1, label, border=True, font=book.italic if label.startswith("—") else book.font)
        book.put(ws, row, 2, number, border=True, font=book.bold)
        book.put(ws, row, 3, share, fmt="0%", border=True, font=book.muted if share is not None else book.font)
        row += 1

    row += 1
    book.put(ws, row, 1, "Operatorlar", font=book.section_font)
    row += 1
    book.header(ws, row, ["Operator", "Username", "Jami"] + _status_titles(report) + ["Asosiy natija"])
    row += 1
    if not report.operators:
        book.put(ws, row, 1, "Bu davrda qo'ng'iroq qayd etilmadi", font=book.muted)
        row += 1
    total = {"total": 0, "positive": 0, "counts": {}}
    for line in report.operators:
        name = line.name if line.listed else f"➕ {line.name}"
        values = [name, f"@{line.username}" if line.username else "", line.total]
        values += [line.counts.get(emoji, 0) for emoji, _ in report.columns]
        values.append(_share(line.positive, line.total))
        book.values(ws, row, values, font=None if line.listed else book.italic, fmts={len(values): "0%"})
        total["total"] += line.total
        total["positive"] += line.positive
        for emoji, count in line.counts.items():
            total["counts"][emoji] = total["counts"].get(emoji, 0) + count
        row += 1
    if report.operators:
        values = ["JAMI", "", total["total"]] + [total["counts"].get(emoji, 0) for emoji, _ in report.columns]
        values.append(_share(total["positive"], total["total"]))
        book.values(ws, row, values, font=book.bold, fill=book.total_fill, fmts={len(values): "0%"})
        row += 1
        if any(not line.listed for line in report.operators):
            book.put(ws, row, 1, "➕ — ro'yxatda yo'q, lekin qo'ng'iroq qilgan", font=book.muted)
            row += 1
        book.put(ws, row, 1, "Bir davrda bitta lid bir operator uchun bir marta sanaladi (so'nggi natija).", font=book.muted)
        row += 1

    if report.courses:
        row += 1
        book.put(ws, row, 1, "Kurslar bo'yicha", font=book.section_font)
        row += 1
        book.header(ws, row, ["Kurs", "Belgilangan lidlar", report.positive_label, "Ulush"])
        row += 1
        for course, count, positive in report.courses:
            book.values(ws, row, [course, count, positive, _share(positive, count)], fmts={4: "0%"})
            row += 1

    if report.reasons:
        row += 1
        book.put(ws, row, 1, "Rad etish sabablari (izohlardan)", font=book.section_font)
        row += 1
        book.header(ws, row, ["Sabab", "Soni"])
        row += 1
        for reason, count in report.reasons:
            book.values(ws, row, [reason, count])
            row += 1


def _days_sheet(book: _Book, report: ExcelReport) -> None:
    cols = len(report.columns)
    ws = book.sheet("Kunlar", [12, 16, 11] + [11] * cols + [12, 11, 14])
    titles = ["Sana", "Kun", "Qo'ng'iroq"] + _status_titles(report) + ["Asosiy natija", "Yangi lidlar", "Belgilanmagan"]
    book.header(ws, 1, titles)
    ws.freeze_panes = "A2"
    row = 2
    total = {"total": 0, "positive": 0, "new": 0, "unmarked": 0, "counts": {}}
    for line in report.days:
        values: list[Any] = [line.day, line.weekday + (" (hozircha)" if line.live else ""), line.total]
        values += [line.counts.get(emoji, 0) for emoji, _ in report.columns]
        values += [_share(line.positive, line.total), line.new_leads, line.unmarked]
        book.values(ws, row, values, fmts={1: "DD.MM.YYYY", 4 + cols: "0%"})
        total["total"] += line.total
        total["positive"] += line.positive
        total["new"] += line.new_leads
        total["unmarked"] += line.unmarked
        for emoji, count in line.counts.items():
            total["counts"][emoji] = total["counts"].get(emoji, 0) + count
        row += 1
    if not report.days:
        book.put(ws, row, 1, "Bu davrda hisobot kuni yo'q", font=book.muted)
        return
    values = ["Jami", "kunlar yig'indisi", total["total"]] + [total["counts"].get(emoji, 0) for emoji, _ in report.columns]
    values += [_share(total["positive"], total["total"]), total["new"], total["unmarked"]]
    book.values(ws, row, values, font=book.bold, fill=book.total_fill, fmts={4 + cols: "0%"})
    book.put(ws, row + 2, 1, "Har kun — o'sha kungi hisobot davri (masalan, kecha 19:00 → bugun 19:00). Bir lidga turli kunlarda "
             "qilingan qo'ng'iroq har kunda alohida sanaladi; «Umumiy» varaqda — butun davr bo'yicha bir marta.", font=book.muted)


def _leads_sheet(book: _Book, report: ExcelReport, masked: bool) -> None:
    ws = book.sheet("Lidlar", [5, 17, 22, 20, 7, 16, 24, 18, 20, 40, 16, 16, 12])
    titles = ["№", "Sana", "Ism", "Telefon", "Yosh", "Tuman", "Kurs", "Natija", "Kim belgiladi", "So'nggi izoh",
              "Sabab", "Qayta qo'ng'iroq", "Havola"]
    book.header(ws, 1, titles)
    ws.freeze_panes = "A2"
    if not report.leads:
        book.put(ws, 2, 1, "Bu davrda yangi lid tushmadi", font=book.muted)
        return
    for index, lead in enumerate(report.leads, 1):
        row = index + 1
        phones = ", ".join((mask_phone if masked else format_phone)(p) for p in lead.phones if p)
        values = [index, lead.created, lead.name, phones, lead.age, lead.district, lead.course, lead.status,
                  lead.people, lead.note[:NOTE_LIMIT], lead.reason, lead.callback, "ochish" if lead.link else ""]
        book.values(ws, row, values, fmts={2: "DD.MM.YYYY HH:MM"}, wrap=True)
        status_cell = ws.cell(row=row, column=8)
        if lead.positive:
            status_cell.fill = book.good_fill
        elif lead.status == "belgilanmagan":
            status_cell.fill = book.warn_fill
        if lead.link:
            link = ws.cell(row=row, column=13)
            link.hyperlink = lead.link
            link.font = book.link_font
    ws.auto_filter.ref = f"A1:M{len(report.leads) + 1}"


def _tasks_sheet(book: _Book, report: ExcelReport) -> None:
    ws = book.sheet("Vazifalar", [5, 20, 46, 16, 16, 18, 16, 10, 22])
    titles = ["№", "Ijrochi", "Vazifa", "Berilgan", "Muddat", "Holat", "Yopilgan", "«Yo'q» javoblari", "Takror"]
    book.header(ws, 1, titles)
    ws.freeze_panes = "A2"
    if not report.tasks:
        book.put(ws, 2, 1, "Bu davrda vazifa yo'q", font=book.muted)
        return
    for index, task in enumerate(report.tasks, 1):
        row = index + 1
        values = [index, task.who, task.text, task.created, task.due, task.state, task.closed, task.no_count or None,
                  task.repeat]
        book.values(ws, row, values, fmts={4: "DD.MM.YYYY HH:MM", 5: "DD.MM.YYYY HH:MM", 7: "DD.MM.YYYY HH:MM"}, wrap=True)
        state = ws.cell(row=row, column=6)
        if task.state.startswith("✅"):
            state.fill = book.good_fill
        elif task.state.startswith(("🔴", "⚠️", "❌")):
            state.fill = book.warn_fill
    ws.auto_filter.ref = f"A1:I{len(report.tasks) + 1}"


def build_workbook(report: ExcelReport, *, masked: bool) -> bytes:
    """Excel fayl (bayt). masked=True — rahbarlar guruhi uchun: telefonlar yashiriladi."""
    if masked:
        report = masked_report(report)
    book = _Book()
    _summary_sheet(book, report, masked)
    _days_sheet(book, report)
    _leads_sheet(book, report, masked)
    _tasks_sheet(book, report)
    return book.save()
