"""v1.9: kundalik (egasiga, ish kunlari 20:00) va haftalik / oylik Excel (rahbarlarga yashirilgan, egasiga to'liq)."""
import io
from datetime import date, datetime, timezone

from aiogram.exceptions import TelegramForbiddenError, TelegramNetworkError

from bot.reports.archive import MONTH, WEEK, due_period, parse_key, period_key, period_label, previous_period
from bot.reports.excel import format_phone, mask_phone
from test_flow import OWNER, run
from test_reports import (
    CARD_1,
    CARD_2,
    LEADERS,
    MALIKA,
    SARDOR,
    WED,
    classify,
    local,
    person,
    setup_groups,
    start,
    ts,
)
from test_reports_api import LEADER, Api

JASUR = 503
MON = date(2026, 10, 19)
TUE = date(2026, 10, 20)
THU = date(2026, 10, 22)
FRI = date(2026, 10, 23)
SAT = date(2026, 10, 24)
NEXT_MON = date(2026, 10, 26)
NOV_2 = date(2026, 11, 2)
COMMENT = "ko'tarmadi, keyin +998 99 888 77 66 ga qayta qilaman"


def workbook(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


def column(ws, title):
    """Sarlavhasi `title` bo'lgan ustun qiymatlari (birinchi qator — sarlavha)."""
    headers = [c.value for c in ws[1]]
    index = headers.index(title)
    return [row[index] for row in ws.iter_rows(min_row=2, values_only=True)]


def find_row(ws, first):
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == first:
            return row
    raise AssertionError(f"{first!r} qatori topilmadi")


async def seed_week(env):
    """19–24 oktabr haftasi: 4 ta lid (biri belgilanmagan), qayta qo'ng'iroq, ro'yxatda yo'q operator, 3 ta vazifa."""
    g = await setup_groups(env)
    reports = env.ctx.reports
    store = reports.store
    await store.add_person("operator", username="malika_op", user_id=None, name="Malika")
    await store.add_person("operator", username="sardor", user_id=None, name="Sardor")
    await env.db.set_config("report:archive_since", "2026-10-01T00:00:00")
    malika, sardor, jasur = person(MALIKA, "Malika", "malika_op"), person(SARDOR, "Sardor", "sardor"), person(JASUR, "Jasur")

    a = await g.post(CARD_1, ts(MON, "10:00"))
    await g.react(a["message_id"], malika, "👍", ts(MON, "10:30"))
    b = await g.post(CARD_2, ts(TUE, "11:00"))
    await g.react(b["message_id"], sardor, "🆒", ts(TUE, "11:20"))
    await g.react(b["message_id"], sardor, "👎", ts(WED, "12:00"))  # ertasi kuni qayta qo'ng'iroq
    c = await g.post("Ism: Dilshod\nTelefon: 99 888 77 66\nKurs: Kompyuter", ts(WED, "09:30"))
    await g.comment(c, malika, COMMENT, ts(WED, "09:45"))
    await g.post("Ism: Nodira\nTelefon: 95 111 22 33\nKurs: Mobilografiya", ts(THU, "10:00"))  # hech kim belgilamadi
    await g.react(a["message_id"], jasur, "🤔", ts(THU, "10:30"))  # ro'yxatda yo'q
    await g.post("Ism: Kecha\nTelefon: 90 000 00 01\nKurs: Dizayn", ts(SAT, "20:00"))  # shanba 19:00 dan keyin — keyingi hafta

    tasks = reports.tasks
    one = await tasks.create(creator_id=OWNER, assignee=await tasks.resolve(MALIKA, "malika_op", "Malika"),
                             text="Oylik hisobotni topshirish", due=local(WED, "18:00"), repeat=None, poll_chat=None)
    two = await tasks.create(creator_id=OWNER, assignee=await tasks.resolve(SARDOR, "sardor", "Sardor"),
                             text="Shartnomani +998 90 555 66 77 bilan kelishish", due=local(THU, "18:00"), repeat=None,
                             poll_chat=None)
    three = await tasks.create(creator_id=OWNER, assignee=await tasks.resolve(MALIKA, "malika_op", "Malika"),
                               text="Lidlar ro'yxatini tekshirish", due=local(date(2026, 10, 30), "18:00"), repeat=None,
                               poll_chat=None)
    def utc(moment):
        return moment.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S")
    await tasks.store.update(one.id, status="done", done_at=utc(local(WED, "17:00")), done_by=MALIKA)
    for task, created in ((one, local(WED, "09:00")), (two, local(WED, "10:00")), (three, local(FRI, "11:00"))):
        await tasks.store.conn.execute("UPDATE rp_tasks SET created_at = ? WHERE id = ?", (utc(created), task.id))
    await tasks.store.conn.commit()
    return g, (a, b, c)


def queue_analysis(env):
    """C lid izohini tahlil qilish javobi (birinchi tahlil chaqiruvi oladi)."""
    async def go():
        pending = await env.ctx.reports.store.pending_comments(10)
        env.api.responses.append(classify(
            {"id": pending[0].id, "call": True, "status": "🆒", "reason": "", "callback": "keyin"}))
    return go()


# --------------------------------------------------------------------------- davrlar
def test_period_keys_and_labels():
    assert period_key(WEEK, MON) == "week:2026-10-19" and period_key(MONTH, date(2026, 10, 1)) == "month:2026-10"
    assert parse_key("week:2026-10-19") == (WEEK, MON, date(2026, 10, 25))
    assert parse_key("month:2026-02") == (MONTH, date(2026, 2, 1), date(2026, 2, 28))
    assert parse_key("month:2026-12") == (MONTH, date(2026, 12, 1), date(2026, 12, 31))
    assert parse_key("week:2026-10-20") is None  # dushanba emas
    assert parse_key("week:2026-13-01") is None and parse_key("year:2026") is None and parse_key("") is None
    assert period_label(WEEK, MON, date(2026, 10, 25)) == "19–25 oktabr 2026"
    assert period_label(WEEK, date(2026, 9, 28), date(2026, 10, 4)) == "28 sentabr – 4 oktabr 2026"
    assert period_label(WEEK, date(2025, 12, 29), date(2026, 1, 4)) == "29 dekabr 2025 – 4 yanvar 2026"
    assert period_label(MONTH, date(2026, 10, 1), date(2026, 10, 31)) == "Oktabr 2026"
    assert due_period(WEEK, NEXT_MON) == (MON, date(2026, 10, 25), NEXT_MON)
    assert due_period(WEEK, date(2026, 10, 25)) == (date(2026, 10, 12), date(2026, 10, 18), MON)  # yakshanba
    assert due_period(MONTH, NOV_2) == (date(2026, 10, 1), date(2026, 10, 31), date(2026, 11, 1))
    assert due_period(MONTH, date(2026, 1, 15)) == (date(2025, 12, 1), date(2025, 12, 31), date(2026, 1, 1))
    assert previous_period(MONTH, date(2026, 3, 1)) == (date(2026, 2, 1), date(2026, 2, 28))
    assert format_phone("+998901234567") == "+998 90 123-45-67" and mask_phone("+998901234567") == "+998 90 ***-**-67"
    assert mask_phone("12345") == "***"


# --------------------------------------------------------------------------- Excel
def test_weekly_excel_masked_and_full(tmp_path):
    async def go():
        env = await start(tmp_path)
        try:
            await seed_week(env)
            await queue_analysis(env)
            archive = env.ctx.reports.archive
            now = local(NEXT_MON, "09:05")
            export = await archive.collect(WEEK, MON, date(2026, 10, 25), now)
            s = export.summary
            assert (s["calls"], s["positive"], s["positive_pct"], s["new_leads"], s["unmarked"]) == (4, 1, 25, 4, 1)
            assert (s["tasks_due"], s["tasks_done"], s["tasks_on_time"], s["tasks_open"]) == (2, 1, 1, 1)
            assert s["top"][0] == ["Malika", 2, 1] and export.period.live is False
            assert export.report.filename == "Haftalik_hisobot_2026-10-19_2026-10-25.xlsx"
            full, masked = await archive.files(export, False, True)

            book = workbook(full)
            assert book.sheetnames == ["Umumiy", "Kunlar", "Lidlar", "Vazifalar"]
            summary = book["Umumiy"]
            assert summary["A1"].value == "Haftalik hisobot · 19–25 oktabr 2026"
            assert summary["A2"].value.startswith("Hisobot davri: 17.10.2026 19:00 → 24.10.2026 19:00 · 6 ish kuni")
            assert find_row(summary, "Qo'ng'iroq qilingan lidlar")[1] == 4
            assert find_row(summary, "Belgilanmagan yangi lidlar")[1:3] == (1, 0.25)
            malika = find_row(summary, "Malika")
            assert malika[1:3] == ("@malika_op", 2) and malika[-1] == 0.5
            assert find_row(summary, "➕ Jasur")[2] == 1  # ro'yxatda yo'q
            assert find_row(summary, "JAMI")[2] == 4

            days = book["Kunlar"]
            assert column(days, "Qo'ng'iroq")[:6] == [1, 1, 2, 1, 0, 0]  # dushanba–shanba
            assert column(days, "Yangi lidlar")[:6] == [1, 1, 1, 1, 0, 0]
            assert find_row(days, "Jami")[2] == 5  # kunlar yig'indisi (B lid ikki kunda)

            leads = book["Lidlar"]
            assert column(leads, "Ism") == ["Aziza Karimova", "Bobur", "Dilshod", "Nodira"]  # shanba 20:00 dagi — yo'q
            assert column(leads, "Telefon")[:2] == ["+998 90 123-45-67", "+998 93 765-43-21, +998 97 111-22-33"]
            assert column(leads, "Natija") == ["🤔 o'ylayapti", "👎 kelmaydi", "🆒 ko'tarmadi", "belgilanmagan"]
            assert column(leads, "Kim belgiladi")[0] == "Malika, Jasur"
            assert column(leads, "So'nggi izoh")[2] == COMMENT and column(leads, "Qayta qo'ng'iroq")[2] == "keyin"
            assert leads.cell(row=2, column=13).hyperlink.target.startswith("https://t.me/c/1111111111/")
            tasks = book["Vazifalar"]
            assert column(tasks, "Vazifa") == ["Oylik hisobotni topshirish", "Shartnomani +998 90 555 66 77 bilan kelishish",
                                                "Lidlar ro'yxatini tekshirish"]
            assert column(tasks, "Holat") == ["✅ bajarildi", "🔴 muddati o'tdi", "⏳ kutilmoqda"]

            hidden = workbook(masked)
            phones = column(hidden["Lidlar"], "Telefon")
            assert phones[:3] == ["+998 90 ***-**-67", "+998 93 ***-**-21, +998 97 ***-**-33", "+998 99 ***-**-66"]
            assert column(hidden["Lidlar"], "So'nggi izoh")[2] == "ko'tarmadi, keyin [raqam] ga qayta qilaman"
            assert column(hidden["Vazifalar"], "Vazifa")[1] == "Shartnomani [raqam] bilan kelishish"
            flat = " ".join(str(v) for ws in hidden for row in ws.iter_rows(values_only=True) for v in row if v)
            for raw in ("123-45-67", "888 77 66", "555 66 77", "765-43-21"):
                assert raw not in flat  # yashirilgan nusxada to'liq raqam hech qayerda yo'q
            assert "Rahbarlar nusxasi" in flat
        finally:
            await env.close()
    run(go())


def test_weekly_and_monthly_delivery(tmp_path):
    async def go():
        env = await start(tmp_path)
        try:
            await seed_week(env)
            await queue_analysis(env)
            archive = env.ctx.reports.archive
            bot = env.bot
            env.session.clear()
            assert await archive.check_exports(bot, local(NEXT_MON, "08:59")) == 0  # hali vaqti emas
            assert await archive.check_exports(bot, local(NEXT_MON, "09:01")) == 1
            docs = env.session.of("SendDocument")
            assert [d.chat_id for d in docs] == [LEADERS, OWNER]
            assert "🔒" in docs[0].caption and "🔓" in docs[1].caption and "19–25 oktabr" in docs[0].caption
            assert "📞 4 ta lidga qo'ng'iroq · 👍 1 (25%)" in docs[1].caption
            group_file, owner_file = env.session.documents[-2][1], env.session.documents[-1][1]
            assert "***" in column(workbook(group_file)["Lidlar"], "Telefon")[0]
            assert column(workbook(owner_file)["Lidlar"], "Telefon")[0] == "+998 90 123-45-67"
            row = await archive.store.export("week:2026-10-19")
            assert row["sent_at"] and '"calls": 4' in row["summary"]
            # guruhdagi faylga reaksiya — qo'ng'iroq emas
            async with env.db.conn.execute("SELECT message_id FROM rp_seen WHERE chat_id = ?", (LEADERS,)) as cur:
                seen = [r[0] for r in await cur.fetchall()]
            assert env.session.next_id - 1 in seen  # guruh nusxasi egasinikidan oldin yuborildi
            assert await archive.check_exports(bot, local(NEXT_MON, "12:00")) == 0  # ikkinchi marta yo'q

            # oylik: 2-noyabr (dushanba) — o'tgan hafta va oktabr birga
            env.session.clear()
            assert await archive.check_exports(bot, local(NOV_2, "09:30")) == 2
            names = [d[0] for d in env.session.documents[-4:]]
            assert names == ["Haftalik_hisobot_2026-10-26_2026-11-01.xlsx"] * 2 + ["Oylik_hisobot_2026-10.xlsx"] * 2
            month = workbook(env.session.documents[-1][1])
            assert month["Umumiy"]["A1"].value == "Oylik hisobot · Oktabr 2026"
            notes = [month["Umumiy"].cell(row=r, column=1).value for r in range(3, 7)]
            assert any(n and n.startswith("Kuzatish 01.10.2026 05:00 dan boshlangan") for n in notes)
            assert sum(isinstance(v, datetime) for v in column(month["Kunlar"], "Sana")) == 27  # oktabr ish kunlari
            assert column(month["Lidlar"], "Ism")[-1] == "Kecha"  # shanba 20:00 — dushanba davri (oktabr ichida)
        finally:
            await env.close()
    run(go())


def test_delivery_rules(tmp_path):
    async def go():
        env = await start(tmp_path)
        try:
            await seed_week(env)
            await queue_analysis(env)
            reports = env.ctx.reports
            archive = reports.archive
            bot = env.bot
            # o'rnatilishidan oldingi davr avtomatik yuborilmaydi
            await env.db.set_config("report:archive_since", "2026-10-26T05:00:00")  # dushanba 10:00 (Toshkent)
            assert await archive.check_exports(bot, local(NEXT_MON, "10:30")) == 0
            await env.db.set_config("report:archive_since", "2026-10-01T00:00:00")
            # haftalik o'chirilgan
            await reports.save_schedule({"weekly_enabled": False})
            assert await archive.check_exports(bot, local(NEXT_MON, "10:30")) == 0
            await reports.save_schedule({"weekly_enabled": True, "excel_time": "11:00"})
            assert await archive.check_exports(bot, local(NEXT_MON, "10:30")) == 0  # vaqt o'zgardi

            # rahbarlar guruhiga yozib bo'lmadi — egasiga yetadi va xabar qilinadi
            original = env.session.make_request

            async def flaky(bot_, method, timeout=None):
                if type(method).__name__ == "SendDocument" and method.chat_id == LEADERS:
                    raise TelegramForbiddenError(method=method, message="Forbidden: bot was kicked")
                return await original(bot_, method, timeout)
            env.session.make_request = flaky
            env.session.clear()
            assert await archive.check_exports(bot, local(NEXT_MON, "11:05")) == 0  # guruhga yetmadi — davr ochiq
            assert [d.chat_id for d in env.session.of("SendDocument")] == [OWNER]  # guruhga urinish xato bilan tugadi
            warn = env.session.of("SendMessage")[-1]
            assert warn.chat_id == OWNER and "rahbarlar guruhiga yuborilmadi" in warn.text and "uzatmang" in warn.text
            row = await archive.store.export("week:2026-10-19")
            assert row["sent_at"] is None and archive.store.recipients(row) == {OWNER}
            assert await archive.check_exports(bot, local(NEXT_MON, "11:14")) == 0  # 10 daqiqa kutiladi
            assert len(env.session.of("SendDocument")) == 1
            assert await archive.check_exports(bot, local(NEXT_MON, "11:16")) == 0  # yana xato → endi 20 daqiqa
            assert [d.chat_id for d in env.session.of("SendDocument")] == [OWNER]  # egasiga qayta yuborilmadi
            assert len([m for m in env.session.of("SendMessage") if "yuborilmadi" in m.text]) == 1  # bir marta aytiladi
            env.session.make_request = original
            env.session.clear()
            assert await archive.check_exports(bot, local(NEXT_MON, "11:30")) == 0  # 11:36 gacha kutiladi
            assert await archive.check_exports(bot, local(NEXT_MON, "11:37")) == 1
            assert [d.chat_id for d in env.session.of("SendDocument")] == [LEADERS]  # faqat yetmagan guruhga
            assert "🔒" in env.session.of("SendDocument")[0].caption
            row = await archive.store.export("week:2026-10-19")
            assert row["sent_at"] and archive.store.recipients(row) == {OWNER, LEADERS}
            assert await archive.check_exports(bot, local(NEXT_MON, "12:00")) == 0

            # hech qayerga yetmasa — 10 daqiqadan keyin qayta
            await env.db.conn.execute("DELETE FROM rp_exports")
            await env.db.conn.commit()

            async def broken(bot_, method, timeout=None):
                if type(method).__name__ == "SendDocument":
                    raise TelegramForbiddenError(method=method, message="Forbidden")
                return await original(bot_, method, timeout)
            env.session.make_request = broken
            assert await archive.check_exports(bot, local(NEXT_MON, "12:00")) == 0
            assert (await archive.store.export("week:2026-10-19"))["sent_at"] is None
            env.session.make_request = original
            assert await archive.check_exports(bot, local(NEXT_MON, "12:05")) == 0  # kutish
            assert await archive.check_exports(bot, local(NEXT_MON, "12:11")) == 1

            # tayyorlashda xato — egasiga bir marta aytiladi, keyinroq qayta
            await env.db.conn.execute("DELETE FROM rp_exports")
            await env.db.conn.commit()

            async def crash(*args, **kwargs):
                raise RuntimeError("sinov")
            archive.collect, collect = crash, archive.collect
            env.session.clear()
            assert await archive.check_exports(bot, local(NEXT_MON, "13:00")) == 0
            assert await archive.check_exports(bot, local(NEXT_MON, "13:11")) == 0
            notes = [m.text for m in env.session.of("SendMessage")]
            assert len(notes) == 1 and "tayyorlanmadi" in notes[0] and "/excel" in notes[0]
            archive.collect = collect
            # bot o'chiq bo'lsa — haftalik 2 kun ichida; undan keyin o'tgan hafta avtomatik kelmaydi
            assert await archive.check_exports(bot, local(date(2026, 10, 28), "11:01")) == 0
            assert await archive.check_exports(bot, local(date(2026, 10, 28), "10:59")) == 1
            # o'tgan hafta yuborilmagan bo'lsa ham, keyingi hafta (dushanbadan) faqat yangi davr ko'riladi
            assert due_period(WEEK, date(2026, 11, 3))[0] == NEXT_MON
        finally:
            await env.close()
    run(go())



def test_warning_waits_for_telegram(tmp_path):
    """Ogohlantirish egasiga yetmasa (aloqa yo'q) — «aytildi» deb hisoblanmaydi, keyingi urinishda yuboriladi."""
    async def go():
        env = await start(tmp_path)
        try:
            await seed_week(env)
            await queue_analysis(env)
            archive = env.ctx.reports.archive
            original = env.session.make_request

            async def offline(bot_, method, timeout=None):
                raise TelegramNetworkError(method=method, message="Request timeout error")

            async def group_refuses(bot_, method, timeout=None):
                if type(method).__name__ == "SendDocument" and method.chat_id == LEADERS:
                    raise TelegramForbiddenError(method=method, message="Forbidden: bot was kicked")
                return await original(bot_, method, timeout)
            env.session.make_request = offline
            assert await archive.check_exports(env.bot, local(NEXT_MON, "09:01")) == 0
            env.session.make_request = group_refuses
            env.session.clear()
            assert await archive.check_exports(env.bot, local(NEXT_MON, "09:12")) == 0
            assert [d.chat_id for d in env.session.of("SendDocument")] == [OWNER]
            assert [m.chat_id for m in env.session.of("SendMessage") if "rahbarlar guruhiga yuborilmadi" in m.text] == [OWNER]
            row = await archive.store.export("week:2026-10-19")
            assert row["sent_at"] is None and archive.store.recipients(row) == {OWNER}
        finally:
            await env.close()
    run(go())

# --------------------------------------------------------------------------- kundalik
def test_journal(tmp_path):
    async def go():
        env = await start(tmp_path)
        try:
            await seed_week(env)
            await queue_analysis(env)
            archive = env.ctx.reports.archive
            bot = env.bot
            env.session.clear()
            assert await archive.check_journal(bot, local(WED, "19:59")) is False
            assert await archive.check_journal(bot, local(WED, "20:00")) is True
            sent = env.session.of("SendMessage")
            assert [m.chat_id for m in sent] == [OWNER]  # faqat egasiga, guruhga emas
            text = sent[0].text
            assert "KUNDALIK · 21.10.2026 (chorshanba)" in text
            assert "20.10 19:00 → 21.10 19:00" in text and "Qo'ng'iroq: <b>2</b> ta lid · 👍 keladi <b>0</b> (0%)" in text
            assert "👎 1 · 🆒 1" in text and "📥 Yangi lidlar: 1 · belgilanmagan: 0" in text
            assert "🏆 Malika 1 (👍 0) · Sardor 1 (👍 0)" in text
            assert "✅ Bugun bajarildi: 1\n• <b>Malika</b> — Oylik hisobotni topshirish" in text
            assert "🆕 Bugun berildi: 2" in text and "ochiq: 2" in text
            assert "⏰ Yaqin muddatlar (22.10 gacha): 1\n• <b>Sardor</b> — Shartnomani" in text
            assert "🔴" not in text  # muddati o'tgan yo'q
            assert (await archive.store.journal("2026-10-21"))["text"] == text
            assert await archive.check_journal(bot, local(WED, "20:30")) is False  # bir kunga bitta

            # yarim tundan keyin yozilgan chorshanba kundaligi: «yaqin muddatlar» chorshanbadan hisoblanadi
            late = await archive.journal_text(local(THU, "00:30"), day=WED)
            assert "KUNDALIK · 21.10.2026 (chorshanba)" in late and "⏰ Yaqin muddatlar (22.10 gacha): 1" in late
            # ertasi: muddati o'tgan vazifa; 4 soatdan ko'p kechiksa — yuborilmaydi
            assert await archive.check_journal(bot, local(date(2026, 10, 23), "00:01")) is False  # juma 20:00 emas
            assert await archive.check_journal(bot, local(THU, "23:59")) is True
            text = env.session.of("SendMessage")[-1].text
            assert "🔴 Muddati o'tgan: 1\n• <b>Sardor</b> — Shartnomani +998 90 555 66 77 bilan kelishish · 22.10 (pa) 18:00" in text
            await env.ctx.reports.save_schedule({"journal_time": "12:00"})
            assert await archive.check_journal(bot, local(FRI, "16:01")) is False  # 4 soatdan ko'p kechikdi
            assert await archive.check_journal(bot, local(FRI, "15:59")) is True
            # yakshanba — ish kuni emas
            assert await archive.check_journal(bot, local(date(2026, 10, 25), "12:30")) is False
            # 22:00 dagi kundalik yarim tundan keyin ham (4 soat ichida) keladi — o'sha kun nomi bilan
            await env.ctx.reports.save_schedule({"journal_time": "22:00"})
            assert await archive.check_journal(bot, local(date(2026, 10, 25), "00:30")) is True
            assert "KUNDALIK · 24.10.2026 (shanba)" in env.session.of("SendMessage")[-1].text
            assert await archive.check_journal(bot, local(date(2026, 10, 25), "01:00")) is False  # bir marta
            # o'chirilgan bo'lsa — yo'q
            await env.ctx.reports.save_schedule({"journal_enabled": False})
            assert await archive.check_journal(bot, local(NEXT_MON, "22:00")) is False
        finally:
            await env.close()
    run(go())


def test_commands(tmp_path):
    async def go():
        env = await start(tmp_path)
        try:
            await seed_week(env)
            await queue_analysis(env)
            env.ctx.reports.now = lambda: local(NEXT_MON, "15:00")
            env.session.clear()
            await env.send_text("/kundalik")
            text = env.session.of("SendMessage")[-1].text
            assert "KUNDALIK · 26.10.2026 (dushanba)" in text and "hozirgi holat · 15:00" in text
            assert "har ish kuni 20:00 da o'zi keladi" in text
            await env.send_text("/excel")
            menu = env.session.of("SendMessage")[-1]
            buttons = [b.callback_data for row in menu.reply_markup.inline_keyboard for b in row if b.callback_data]
            assert buttons == ["xl:week:2026-10-19", "xl:week:2026-10-26", "xl:month:2026-09", "xl:month:2026-10"]
            assert "har dushanba 09:00 da" in menu.text
            await env.press("xl:week:2026-10-19")
            name, data, caption = env.session.documents[-1]
            assert name == "Haftalik_hisobot_2026-10-19_2026-10-25.xlsx" and "🔓" in caption
            assert column(workbook(data)["Lidlar"], "Telefon")[0] == "+998 90 123-45-67"
            await env.press("xl:week:2026-10-26")  # joriy hafta — oraliq
            assert env.session.documents[-1][0] == "Haftalik_hisobot_2026-10-26_2026-11-01_oraliq.xlsx"
            assert (await env.ctx.reports.archive.store.export("week:2026-10-26")) is None  # joriy davr saqlanmaydi
            await env.press("xl:week:2026-10-20")
            assert "eskirgan" in env.session.of("AnswerCallbackQuery")[-1].text
            # egasi bo'lmagan odam
            await env.db.add_member(777, "Ali", "ali")
            await env.press("xl:week:2026-10-19", user_id=777)
            assert "Faqat bot egasi" in env.session.of("AnswerCallbackQuery")[-1].text
        finally:
            await env.close()
    run(go())


# --------------------------------------------------------------------------- Mini App API
def test_archive_api(tmp_path):
    async def go():
        api = await Api().start(tmp_path)
        try:
            env = api.env
            await seed_week(env)
            await queue_analysis(env)
            reports = env.ctx.reports
            reports.now = lambda: local(NEXT_MON, "15:00")
            await reports.store.add_person("leader", username="rahbar_1", user_id=None, name="Rahbar")
            leader = (LEADER, "rahbar_1")
            await reports.archive.check_exports(env.bot, local(NEXT_MON, "09:05"))
            await reports.archive.check_journal(env.bot, local(NEXT_MON, "20:00"))

            weekly = await api.call("GET", "/api/report/archive?kind=weekly")
            items = weekly["items"]
            assert [i["key"] for i in items[:3]] == ["week:2026-10-26", "week:2026-10-19", "week:2026-10-12"]
            assert items[0]["live"] and items[0]["summary"] is None
            assert items[1]["sent"] and items[1]["summary"]["calls"] == 4 and items[1]["label"] == "19–25 oktabr 2026"
            assert items[-1]["key"] == "week:2026-09-28"  # kuzatish 1-oktabrdan: 28.09–04.10 haftasi
            assert weekly["schedule"]["excel_time"] == "09:00" and weekly["leaders_group"] is True
            monthly = await api.call("GET", "/api/report/archive?kind=monthly", user=leader)
            assert [i["key"] for i in monthly["items"]] == ["month:2026-10"]
            daily = await api.call("GET", "/api/report/archive?kind=daily")
            assert daily["items"] == []  # kunlik hisobotlar yuborilmagan (testda)
            journal = await api.call("GET", "/api/report/archive?kind=journal")
            assert [i["key"] for i in journal["items"]] == ["journal:2026-10-26"]
            await api.call("GET", "/api/report/archive?kind=journal", user=leader, status=403)
            detail = await api.call("GET", "/api/report/archive/journal:2026-10-26")
            assert "KUNDALIK" in detail["text"] and "<b>" not in detail["text"]
            await api.call("GET", "/api/report/archive/journal:2026-10-26", user=leader, status=403)
            await api.call("GET", "/api/report/archive/journal:2026-01-01", status=404)

            for key in ("week:0001-01-01", "week:9999-12-27", "week:2026-09-21", "week:2026-11-02", "month:2026-09"):
                await api.call("GET", f"/api/report/archive/{key}", user=leader, status=404)  # kuzatish oralig'idan tashqari
            week = await api.call("GET", "/api/report/archive/week:2026-10-12", user=leader)
            assert week["summary"]["calls"] == 0 and week["masked"] is True and week["sent"] is False
            assert (await reports.archive.store.export("week:2026-10-12"))["sent_at"] is None  # ko'rildi, yuborilmadi
            await api.call("GET", "/api/report/archive/week:2026-10-13", status=404)

            # Excel'ni o'z chatiga: rahbarga yashirilgan, egasiga to'liq
            res = await api.call("POST", "/api/report/archive/week:2026-10-19/send", {}, user=leader)
            assert res == {"ok": True, "masked": True}
            again = await api.call("POST", "/api/report/archive/week:2026-10-19/send", {}, user=leader, status=429)
            assert "hozirgina" in again["message"]
            # fayl yasalmasa — 500, cheklov qo'yilmaydi (keyingi urinish darrov mumkin)
            archive = reports.archive
            collect = archive.collect

            async def crash(*args, **kwargs):
                raise RuntimeError("sinov")
            archive.collect = crash
            failed = await api.call("POST", "/api/report/archive/month:2026-10/send", {}, status=500)
            assert failed["error"] == "build_failed"
            archive.collect = collect
            doc = env.session.of("SendDocument")[-1]
            assert doc.chat_id == LEADER and "🔒" in doc.caption
            assert "***" in column(workbook(env.session.documents[-1][1])["Lidlar"], "Telefon")[0]
            res = await api.call("POST", "/api/report/archive/month:2026-10/send", {})
            assert res["masked"] is False and env.session.of("SendDocument")[-1].chat_id == OWNER
            await api.call("POST", "/api/report/archive/week:2026-10-19/xyz", {}, user=leader, status=403)
            await api.call("POST", "/api/report/archive/week:2026-10-19/send", {}, user=(999, "begona"), status=403)
            denied = await api.call("PATCH", "/api/report/schedule", {"journal_time": "21:00"}, user=leader, status=403)
            assert "egasi" in denied["message"]

            # sozlamalar: yangi maydonlar
            res = await api.call("PATCH", "/api/report/schedule", {"journal_time": "21:30", "weekly_enabled": False,
                                                                    "excel_time": "8:15", "monthly_enabled": False})
            sc = res["schedule"]
            assert (sc["journal_time"], sc["weekly_enabled"], sc["monthly_enabled"], sc["excel_time"]) == (
                "21:30", False, False, "08:15")
            settings = await api.call("GET", "/api/report/settings")
            assert settings["schedule"]["journal_enabled"] is True and settings["schedule"]["journal_time"] == "21:30"
        finally:
            await api.close()
    run(go())


def test_excel_hardening(tmp_path):
    """Yashirilgan nusxada hech qanday maydonda to'liq raqam yo'q; «=…» formula emas; boshqaruv belgisi faylni buzmaydi;
    biror kunda aniqlanmagan natija bo'lsa — «Kunlar»da ❔ ustuni bor."""
    async def go():
        env = await start(tmp_path)
        try:
            g = await setup_groups(env)
            reports = env.ctx.reports
            await reports.store.add_person("operator", username="malika_op", user_id=None, name="Malika")
            malika, sardor = person(MALIKA, "Malika", "malika_op"), person(SARDOR, "Sardor", "sardor")
            lola = await g.post("Ism: Lola (onasi 91 765 43 21)\nTuman: Chilonzor 9, qo'shimcha 93 111 22 33\nKurs: =1+2",
                                ts(MON, "10:00"))
            odd = await g.post("Ism: Ali\uffff Valiyev, ukasi 97 / 111 / 44 / 55\nTelefon: +998 93\u2011765\u201143\u201121"
                               "\nKurs: Excel", ts(MON, "12:00"))
            evil = await g.post('Ism: =HYPERLINK("http://evil.example/?x","Aziza")\nTelefon: 90 222 33 44\nKurs: Excel',
                                ts(MON, "11:00"))
            await g.comment(lola, malika, "qildim", ts(MON, "10:05"))  # natijasi aniqlanmadi → dushanba ❔
            await g.react(lola["message_id"], malika, "👎", ts(TUE, "10:00"))  # haftada — 👎
            await g.comment(evil, sardor, "=SUM(1,2) 90 12 34 567 ga yozdim", ts(MON, "11:30"))
            await g.comment(odd, sardor, "22.10.2026 18.00 da keladi, onasi +998 90–123–45–67", ts(MON, "12:30"))
            pending = await reports.store.pending_comments(10)
            env.api.responses.append(classify(
                {"id": pending[0].id, "call": True, "status": "", "reason": "", "callback": ""},
                {"id": pending[1].id, "call": True, "status": "👎", "reason": "", "callback": ""},
                {"id": pending[2].id, "call": True, "status": "👍", "reason": "", "callback": "22.10.2026 18.00"}))
            tasks = reports.tasks
            await tasks.create(creator_id=OWNER, assignee=await tasks.resolve(MALIKA, "malika_op", "Malika"),
                               text="Hisobot\u0007 tayyorlash", due=local(THU, "18:00"), repeat=None, poll_chat=None)
            export = await reports.archive.collect(WEEK, MON, date(2026, 10, 25), local(NEXT_MON, "09:05"))
            full, masked = await reports.archive.files(export, False, True)

            book = workbook(full)
            names = book["Lidlar"]
            assert names.cell(row=3, column=3).value.startswith("=HYPERLINK") and names.cell(row=3, column=3).data_type == "s"
            assert names.cell(row=3, column=10).value.startswith("=SUM") and names.cell(row=3, column=10).data_type == "s"
            course = find_row(book["Umumiy"], "=1+2")
            assert course[1] == 1
            assert column(book["Vazifalar"], "Vazifa") == ["Hisobot tayyorlash"]
            days = book["Kunlar"]
            headers = [c.value for c in days[1]]
            assert "❔ aniqlanmagan" in headers and "❔ aniqlanmagan" in [c.value for c in book["Umumiy"][find_header(book["Umumiy"])]]
            assert days.cell(row=2, column=headers.index("❔ aniqlanmagan") + 1).value == 1  # dushanba
            assert days.cell(row=3, column=headers.index("👎 kelmaydi") + 1).value == 1  # seshanba

            hidden = workbook(masked)
            flat = " ".join(str(v) for ws in hidden for row in ws.iter_rows(values_only=True) for v in row if v)
            for raw in ("765 43 21", "111 22 33", "222 33 44", "12 34 567", "111 / 44 / 55", "765\u201143", "123–45–67"):
                assert raw not in flat
            assert "Lola (onasi [raqam])" in flat and "qo'shimcha [raqam]" in flat and "[raqam] ga yozdim" in flat
            assert "Ali Valiyev, ukasi [raqam]" in flat  # U+FFFF olib tashlandi, fayl yasaldi
            assert "22.10.2026 18.00 da keladi, onasi [raqam]" in flat  # sana va soat joyida
            assert column(hidden["Lidlar"], "Qayta qo'ng'iroq")[2] == "22.10.2026 18.00"
        finally:
            await env.close()
    run(go())


def find_header(ws):
    """Operatorlar jadvali sarlavhasi qatori raqami."""
    for row in ws.iter_rows():
        if row[0].value == "Operator":
            return row[0].row
    raise AssertionError("Operatorlar jadvali topilmadi")
