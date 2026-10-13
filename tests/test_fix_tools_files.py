"""tools-files tuzatishlari: Word/Excel/PDF/grafik yaratuvchilar, fayl tanlash, hujjat tarjimasi."""
import asyncio
import base64
import io
import zipfile

import docx
import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

import fakes
from bot import doctrans, export, files, media, pdfs, tools
from bot.tools import TEXT_KINDS, RawFile, RawStore, ToolContext
from test_flow import Env, png, run


def tool_ctx(items: list[RawFile]) -> ToolContext:
    return ToolContext(bot=None, chat_id=1, reply_to=1, uid=1, files=items, model="claude-opus-5-5")


def docx_paragraphs(data: bytes) -> list:
    return docx.Document(io.BytesIO(data)).paragraphs


# ------------------------------------------------------------------ 1. Word: _ va * so'z ichida tegilmaydi
def test_docx_export_keeps_underscores_and_asterisks():
    text = "Sozlama: BOT_TOKEN va OWNER_IDS https://site.uz/my_page_name, hisob: 2 * 3 * 4 = 24"
    paragraph = docx_paragraphs(export.answer_to_docx(text))[0]
    assert paragraph.text == text
    assert not any(run.italic for run in paragraph.runs)

    styled = docx_paragraphs(export.answer_to_docx("*kursiv* _ham_ **qalin** __yana__ `a_b_c`"))[0]
    assert styled.text == "kursiv ham qalin yana a_b_c"
    assert [r.text for r in styled.runs if r.italic] == ["kursiv", "ham"]
    assert [r.text for r in styled.runs if r.bold] == ["qalin", "yana"]

    # PDF/pptx matni: kod ichidagi va so'z o'rtasidagi __ saqlanadi
    assert files._plain("Python: `__init__.py`") == "Python: __init__.py"
    assert files._plain("MAX__LEN va __qalin__") == "MAX__LEN va qalin"


# ------------------------------------------------------------------ 2. Word: raqamli ro'yxatlar qayta boshlanadi
def test_docx_numbered_lists_keep_author_numbers():
    data = export.answer_to_docx("## Afzalliklar\n1. Tez\n2. Arzon\n## Kamchiliklar\n1. Murakkab\n2. Sekin")
    paragraphs = docx_paragraphs(data)
    texts = [p.text for p in paragraphs]
    assert texts[1:3] == ["1.\tTez", "2.\tArzon"] and texts[4:6] == ["1.\tMurakkab", "2.\tSekin"]
    assert all(p.style.name != "List Number" for p in paragraphs)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert "ListNumber" not in archive.read("word/document.xml").decode()
    nested = docx_paragraphs(export.answer_to_docx("5. Beshinchi\n   1. ichki"))
    assert nested[0].text == "5.\tBeshinchi"
    assert nested[1].paragraph_format.left_indent > nested[0].paragraph_format.left_indent


# ------------------------------------------------------------------ 3. Excel: ID va ming ajratuvchisi buzilmaydi
def test_xlsx_keeps_ids_and_grouped_numbers():
    row = ["25,000", "1.500", "007", "20208000900123456001", 8600123412341234, "12,5", "3.75", "100", "-0", True]
    sheet = load_workbook(io.BytesIO(files.build_xlsx([{"rows": [row]}]))).active
    values = [cell.value for cell in sheet[1]]
    assert values[:5] == ["25,000", "1.500", "007", "20208000900123456001", "8600123412341234"]
    assert values[5:] == [12.5, 3.75, 100, 0, True]
    # vergulli o'nli sonlar (ming ajratuvchisi bo'la olmaydiganlari) — son
    sheet = load_workbook(io.BytesIO(files.build_xlsx([{"rows": [["0,125", "12,3456", "1234,567", "1,500"]]}]))).active
    assert [cell.value for cell in sheet[1]] == [0.125, 12.3456, 1234.567, "1,500"]


# ------------------------------------------------------------------ 4. Bir xil nomli (stem) fayllar
def test_pick_prefers_exact_name_over_same_stem():
    word = RawFile("shartnoma.docx", "docx", files.build_docx("Narx: 100 so'm"))
    pdf = RawFile("shartnoma.pdf", "pdf", files.build_pdf("Narx: 120 so'm"))
    items = [word, pdf]
    chosen = tools.pick(items, ["shartnoma.docx", "shartnoma.pdf"], TEXT_KINDS, at_least=2)
    assert [f.name for f in chosen] == ["shartnoma.docx", "shartnoma.pdf"]
    assert tools.pick(items, ["shartnoma.docx"], ("docx",)) == [word]
    pdf_first = [RawFile("hisobot.pdf", "pdf", b"%PDF"), RawFile("hisobot.docx", "docx", b"")]
    assert tools.pick(pdf_first, ["hisobot.pdf"], ("pdf",))[0] is pdf_first[0]
    # bitta fayl ataylab ikki marta (merge a.pdf + a.pdf) hali ham ishlaydi
    single = [RawFile("a.pdf", "pdf", b"%PDF"), RawFile("b.pdf", "pdf", b"%PDF")]
    assert [f.name for f in tools.pick(single, ["a.pdf", "a.pdf"], ("pdf",), at_least=2)] == ["a.pdf", "a.pdf"]

    result = run(tools.run(tool_ctx(items), "compare_documents", {"file_a": "shartnoma.docx", "file_b": "shartnoma.pdf"}))
    assert "bir xil" not in result and "-Narx: 100 so'm" in result and "+Narx: 120 so'm" in result

    # Bir xil nomli ikki nusxa: A — eskisi, B — yangisi
    copies = [RawFile("shartnoma.txt", "text", b"Narx: 100\n"), RawFile("shartnoma.txt", "text", b"Narx: 120\n")]
    result = run(tools.run(tool_ctx(copies), "compare_documents", {"file_a": "shartnoma.txt", "file_b": "shartnoma.txt"}))
    assert "-Narx: 100" in result and "+Narx: 120" in result


# ------------------------------------------------------------------ 5. images_to_pdf: oldingi rasmlar aralashmaydi
def test_raw_store_default_is_latest_batch():
    store = RawStore()
    old = RawFile("rasm-1.jpg", "image", b"1")
    store.add(1, [old])
    store.recent(1)  # «bu nima?» navbati
    store.close_batch(1)  # navbat saqlandi
    store.add(1, [RawFile("rasm-3.jpg", "image", b"3")])  # albom yoki alohida yuborilgan,
    store.add(1, [RawFile("rasm-4.jpg", "image", b"4")])  # lekin bitta navbatgacha — bitta to'plam
    current = store.recent(1)
    assert [f.name for f in tools.pick(current, None, ("image",))] == ["rasm-3.jpg", "rasm-4.jpg"]
    assert [f.name for f in tools.pick(current, ["rasm-1.jpg", "rasm-4.jpg"], ("image",))] == ["rasm-1.jpg", "rasm-4.jpg"]
    # oxirgi to'plamda yetarli fayl bo'lmasa — hammasi (merge kamida 2 ta)
    pdf_store = RawStore()
    pdf_store.add(1, [RawFile("a.pdf", "pdf", b"%PDF")])
    pdf_store.close_batch(1)
    pdf_store.add(1, [RawFile("b.pdf", "pdf", b"%PDF")])
    assert [f.name for f in tools.pick(pdf_store.recent(1), None, ("pdf",), at_least=2)] == ["a.pdf", "b.pdf"]


def test_images_to_pdf_ignores_photos_before_new(tmp_path):
    async def go():
        env = await Env().start(tmp_path, [
            fakes.simple_answer("Bu qizil rasm."),
            fakes.tool_answer("pdf_tool", {"action": "images_to_pdf"}),
            fakes.simple_answer("PDF tayyor."),
        ])
        try:
            await env.send_photo("old", png((640, 480)), caption="bu nima?")
            await env.send_text("/new")
            await asyncio.gather(
                env.send_photo("s1", png((800, 1100)), media_group_id="g1"),
                env.send_photo("s2", png((800, 1100)), media_group_id="g1"),
            )
            await env.send_text("PDF qil")
            name, data, _ = env.session.documents[-1]
            assert name == "skaner.pdf" and len(PdfReader(io.BytesIO(data)).pages) == 2
            result = env.api.requests[2]["messages"][-1]["content"][0]["content"]
            assert result.startswith("2 ta rasmdan PDF yuborildi") and "1 ta fayl qo'shilmadi" in result
        finally:
            await env.close()
    run(go())


# ------------------------------------------------------------------ 6. Tarjima: bo'yiga birlashgan katak bir marta
def test_doctrans_vertical_merge_translated_once():
    document = docx.Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(2, 0)).text = "Umumiy shartlar"
    for row, text in enumerate(["Bir", "Ikki", "Uch"]):
        table.cell(row, 1).text = text
    buffer = io.BytesIO()
    document.save(buffer)
    source = buffer.getvalue()
    assert [p.text for p in doctrans.collect(docx.Document(io.BytesIO(source)))] == ["Umumiy shartlar", "Bir", "Ikki", "Uch"]

    seen: list[str] = []

    async def translate(batch):
        seen.extend(batch)
        return [f"[ru] {t}" for t in batch]

    data, count = run(doctrans.translate_docx(source, translate))
    assert count == 4 and seen == ["Umumiy shartlar", "Bir", "Ikki", "Uch"]
    assert docx.Document(io.BytesIO(data)).tables[0].cell(2, 0).text == "[ru] Umumiy shartlar"


# ------------------------------------------------------------------ 7. Grafik: «$» oddiy matn, xatoda rasm yopiladi
def test_chart_dollar_text_and_no_figure_leak(monkeypatch):
    import matplotlib.figure
    import matplotlib.pyplot as plt

    captured = []
    close = plt.close
    monkeypatch.setattr(plt, "close", lambda fig=None: (captured.append(fig), close(fig)))
    spec = {"type": "bar", "title": "Daromad ($) va xarajat ($)", "labels": ["a", "b"],
            "series": [{"name": "Narx ($) va soni ($)", "values": [1, 2]}], "y_label": "$\\frac{a$"}
    assert files.build_chart(spec)[:4] == b"\x89PNG"
    title = captured[-1].axes[0].title
    assert title.get_text() == "Daromad ($) va xarajat ($)" and title._preprocess_math(title.get_text())[1] is False

    before = list(plt.get_fignums())

    def broken(*args, **kwargs):
        raise RuntimeError("disk to'la")

    monkeypatch.setattr(matplotlib.figure.Figure, "savefig", broken)
    for _ in range(3):
        with pytest.raises(RuntimeError):
            files.build_chart({"type": "line", "labels": ["a", "b"], "series": [{"values": [1, 2]}]})
    assert plt.get_fignums() == before
    with pytest.raises(files.FileError, match="manfiy"):
        files.build_chart({"type": "pie", "labels": ["a", "b"], "series": [{"values": [1, -2]}]})


# ------------------------------------------------------------------ 8. Grafik: katta summalar ilmiy ko'rinishsiz
def test_chart_money_values_without_scientific_notation(monkeypatch):
    import matplotlib.pyplot as plt

    captured = []
    close = plt.close
    monkeypatch.setattr(plt, "close", lambda fig=None: (captured.append(fig), close(fig)))
    files.build_chart({"type": "bar", "labels": ["Yan", "Fev", "Mar"],
                       "series": [{"name": "Savdo", "values": [15000000, 23500000, 1234.5]}]})
    ax = captured[-1].axes[0]
    assert [t.get_text() for t in ax.texts] == ["15 000 000", "23 500 000", "1 234.5"]
    assert ax.yaxis.get_major_formatter()(20000000.0) == "20 000 000"
    files.build_chart({"type": "barh", "labels": ["a"], "series": [{"values": [18250000]}]})
    assert captured[-1].axes[0].xaxis.get_major_formatter()(18250000.0) == "18 250 000"
    # 1e10 dan katta summalar (yillik daromad) ham butun yoziladi, o'q bir xil ko'rinishda
    files.build_chart({"type": "bar", "labels": ["2023", "2024", "2025"],
                       "series": [{"name": "Daromad", "values": [12_500_000_000, 23_000_000_000, 18_400_000_000]}]})
    ax = captured[-1].axes[0]
    assert [t.get_text() for t in ax.texts] == ["12 500 000 000", "23 000 000 000", "18 400 000 000"]
    fmt = ax.yaxis.get_major_formatter()
    assert fmt(25_000_000_000.0) == "25 000 000 000"
    assert not any("e" in fmt(v) for v in ax.get_yticks())
    # suzuvchi nuqta shovqini ko'rinmaydi
    assert fmt(0.1 + 0.2) == "0.3" and fmt(0.07500000000000001) == "0.075"


# ------------------------------------------------------------------ 9. CSV ko'p varaq / pptx 12+ band
def test_csv_rejects_extra_sheets_and_pptx_keeps_all_bullets():
    sheets = [{"name": "Yanvar", "columns": ["a"], "rows": [[1]]}, {"name": "Fevral", "columns": ["a"], "rows": [[2]]}]
    with pytest.raises(files.FileError, match="bitta jadval"):
        files.build_csv(sheets)
    with pytest.raises(tools.ToolError, match="bitta jadval"):
        run(tools.run(tool_ctx([]), "create_document", {"format": "csv", "filename": "oylik", "sheets": sheets}))
    assert files.build_csv(sheets[:1]).decode("utf-8-sig") == "a\r\n1\r\n"

    deck = Presentation(io.BytesIO(files.build_pptx(
        [{"title": "Kirish", "bullets": [f"band {i}" for i in range(1, 16)], "notes": "izoh"}]
    )))
    assert len(deck.slides) == 2
    texts = [[s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text] for slide in deck.slides]
    assert texts[0][0] == "Kirish" and texts[1][0] == "Kirish (davomi)"
    bullets = [line for slide in texts for line in slide[1].split("\n")]
    assert len(bullets) == 15 and bullets[-1] == "•  band 15"
    assert deck.slides[0].notes_slide.notes_text_frame.text == "izoh"


# ------------------------------------------------------------------ 10. PDF: tab, emoji, xitoycha
def test_pdf_keeps_tabs_maps_emoji_and_rejects_missing_letters(monkeypatch):
    from fpdf import FPDF

    seen: list[str] = []
    original = FPDF.multi_cell

    def spy(self, w, h=None, text="", *args, **kwargs):
        seen.append(text)
        return original(self, w, h, text, *args, **kwargs)

    monkeypatch.setattr(FPDF, "multi_cell", spy)
    data = files.build_pdf("```go\nfunc main() {\n\tif x {\n```\n\n✅ Bajarildi 🚀\n\n| Holat |\n|---|\n| ❌ |")
    assert "    if x {" in seen and "✔ Bajarildi " in seen
    assert "✘" in PdfReader(io.BytesIO(data)).pages[0].extract_text()
    with pytest.raises(files.FileError, match="docx"):
        files.build_pdf("你好世界\n\n这是一个测试")


# ------------------------------------------------------------------ 11. Rasm → PDF: shaffof fon oq, 16 bit kulrang
def test_images_to_pdf_flattens_alpha_and_scales_16bit():
    from PIL import Image, ImageDraw

    signature = Image.new("RGBA", (200, 100), (0, 0, 0, 0))
    ImageDraw.Draw(signature).line((10, 50, 190, 50), fill=(0, 0, 0, 255), width=6)
    gray16 = Image.new("I;16", (200, 100), 30000)  # ≈ 46% kulrang
    blobs = []
    for img in (signature, gray16):
        buffer = io.BytesIO()
        img.save(buffer, "PNG")
        blobs.append(buffer.getvalue())

    reader = PdfReader(io.BytesIO(pdfs.images_to_pdf(blobs)))
    first = reader.pages[0].images[0].image.convert("RGB")
    assert min(first.getpixel((5, 5))) > 240 and max(first.getpixel((100, 50))) < 40
    second = reader.pages[1].images[0].image.convert("L")
    assert 105 <= second.getpixel((100, 50)) <= 130

    block = media.image_block(blobs[1], 1568)
    sent = Image.open(io.BytesIO(base64.b64decode(block["source"]["data"]))).convert("L")
    assert 105 <= sent.getpixel((100, 50)) <= 130


# ------------------------------------------------------------------ 12. compare_documents: bitta nom berilsa
def test_compare_documents_with_one_name():
    old = RawFile("eski.txt", "text", b"Narx: 100\nMuddat: 10 kun\n")
    new = RawFile("yangi.txt", "text", b"Narx: 120\nMuddat: 10 kun\n")
    for args in ({"file_b": "yangi.txt"}, {"file_a": "eski.txt"}):
        result = run(tools.run(tool_ctx([old, new]), "compare_documents", args))
        assert "A = «eski.txt»" in result and "B = «yangi.txt»" in result and "+Narx: 120" in result
    with pytest.raises(tools.ToolError, match="ikkinchi hujjat yo'q"):
        run(tools.run(tool_ctx([new]), "compare_documents", {"file_b": "yangi.txt"}))
