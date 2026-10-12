"""Vositalar tizimi (v1.4): fayl yaratish, grafik, PDF amallari, hujjat tarjimasi, solishtirish."""
import io
import json
import time
import zipfile

import docx
import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

import fakes
from bot import doctrans, files, pdfs, tools
from bot.claude import MAX_TOOL_ROUNDS
from bot.tools import RawFile, RawStore
from test_flow import OWNER, Env, run


def make_docx(paragraphs: list[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_create_document_flow(tmp_path):
    async def go():
        env = await Env().start(tmp_path, [
            fakes.tool_answer(
                "create_document",
                {"format": "docx", "filename": "Hisobot 2026", "title": "Hisobot",
                 "content": "# Kirish\n\nMatn **qalin**\n\n| Oy | Summa |\n|---|---|\n| Yanvar | 100 |"},
                before="Hisobotni tayyorlayapman.",
            ),
            fakes.simple_answer("📄 Hisobot tayyor — faylda kirish va jadval bor."),
        ])
        try:
            await env.send_text("Hisobotni Word qilib ber")
            first, second = env.api.requests
            names = [t.get("name") for t in first["tools"]]
            assert names == ["web_search", "create_document", "create_chart", "pdf_tool", "translate_document", "compare_documents"]
            assert "Vositalar (fayl bilan ishlash)" in first["system"]
            # ikkinchi so'rov: assistant(tool_use) + user(tool_result)
            assert second["messages"][0] == first["messages"][0]
            assert second["messages"][1]["role"] == "assistant"
            assert [b["type"] for b in second["messages"][1]["content"]] == ["text", "tool_use"]
            assert second["messages"][1]["content"][1]["input"]["filename"] == "Hisobot 2026"
            result = second["messages"][2]
            assert result["role"] == "user" and result["content"][0]["type"] == "tool_result"
            assert result["content"][0]["tool_use_id"] == "toolu_1"
            assert "Fayl yuborildi: Hisobot 2026.docx" in result["content"][0]["content"]
            assert not result["content"][0].get("is_error")

            name, data, caption = env.session.documents[0]
            assert name == "Hisobot 2026.docx" and data[:2] == b"PK" and caption == "📎 Hisobot 2026.docx"
            sent = env.session.of("SendDocument")[0]
            assert sent.reply_parameters.message_id == 101  # foydalanuvchi xabariga javoban
            text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
            assert "Hisobot" in text and "Kirish" in text and "qalin" in text

            final = env.last_edit()
            assert "Hisobotni tayyorlayapman." in final.text and "Hisobot tayyor" in final.text
            assert "📎 1" in final.text and "Opus 5.5" in final.text
            hist = await env.db.history(OWNER, 1, 0)
            assert [m.role for m in hist] == ["user", "assistant"]  # tool bloklari tarixga yozilmaydi
            assert "Hisobot tayyor" in hist[1].content[0]["text"]
            usage = await env.db.spend(OWNER, env.ctx.now().replace(hour=0, minute=0, second=0, microsecond=0))
            assert usage[0] > 0
        finally:
            await env.close()
    run(go())


def test_tool_error_goes_back_to_claude(tmp_path):
    async def go():
        env = await Env().start(tmp_path, [
            fakes.tool_answer("create_document", {"format": "docx", "filename": "x"}),  # content yo'q
            fakes.simple_answer("Kechirasiz, hujjat matnini kiritmabman — mana qayta."),
        ])
        try:
            await env.send_text("Word qil")
            result = env.api.requests[1]["messages"][2]["content"][0]
            assert result["is_error"] is True and "content" in result["content"]
            assert env.session.documents == []
            assert "qayta" in env.last_edit().text and "📎" not in env.last_edit().text.split("<i>")[-1]
        finally:
            await env.close()
    run(go())


def test_chart_and_empty_final_text(tmp_path):
    async def go():
        env = await Env().start(tmp_path, [
            fakes.tool_answer("create_chart", {
                "type": "bar", "title": "Sotuv", "labels": ["Yan", "Fev"],
                "series": [{"name": "2026", "values": [10, 20]}],
            }),
            (200, fakes.sse([fakes.msg_start(), fakes.delta("end_turn", 1), fakes.STOP])),  # matnsiz yakun
        ])
        try:
            await env.send_text("Grafik chiz")
            photo = env.session.of("SendPhoto")[-1]
            assert photo.caption == "Sotuv" and photo.photo.data[:4] == b"\x89PNG"
            assert env.session.photo_uploads == 1
            final = env.last_edit()
            assert "📎 Fayl yuborildi: Sotuv.png" in final.text
            hist = await env.db.history(OWNER, 1, 0)
            assert hist[1].content[0]["text"] == "📎 Fayl yuborildi: Sotuv.png"
        finally:
            await env.close()
    run(go())


def test_pdf_tool_uses_recent_files(tmp_path):
    async def go():
        pdf_a = files.build_pdf("Birinchi hujjat", "A")
        pdf_b = files.build_pdf("Ikkinchi hujjat", "B")
        env = await Env().start(tmp_path, [
            fakes.tool_answer("pdf_tool", {"action": "merge"}),
            fakes.simple_answer("Ikkita PDF birlashtirildi."),
            fakes.tool_answer("pdf_tool", {"action": "extract", "files": ["b.pdf"], "pages": "1"}),
            fakes.simple_answer("1-sahifa olindi."),
        ])
        try:
            await env.send_document("fa", pdf_a, "a.pdf", "application/pdf")
            await env.send_document("fb", pdf_b, "b.pdf", "application/pdf")
            assert "2 ta fayl" in env.session.last_out()
            await env.send_text("ikkalasini birlashtir")
            name, data, _ = env.session.documents[-1]
            assert name == "birlashtirilgan.pdf" and len(PdfReader(io.BytesIO(data)).pages) == 2
            tool_result = env.api.requests[1]["messages"][-1]["content"][0]["content"]
            assert "2 ta PDF birlashtirildi" in tool_result and "2 sahifa" in tool_result
            # fayllar keyingi so'rovda ham mavjud (asl nusxa saqlanadi)
            await env.send_text("endi b.pdf dan 1-sahifani ol")
            name, data, _ = env.session.documents[-1]
            assert name == "b_sahifalar.pdf" and len(PdfReader(io.BytesIO(data)).pages) == 1
            assert "1-sahifa olindi" in env.last_edit().text
        finally:
            await env.close()
    run(go())


def test_translate_document_tool(tmp_path):
    async def go():
        source = make_docx(["Shartnoma", "Tomonlar kelishib oldilar."])
        env = await Env().start(tmp_path, [
            fakes.tool_answer("translate_document", {"target_language": "rus"}),
            fakes.json_answer(json.dumps(["Договор", "Стороны договорились."], ensure_ascii=False), model="claude-opus-5-5", output_tokens=40),
            fakes.simple_answer("Tarjima tayyor."),
        ])
        try:
            await env.send_document("fd", source, "shartnoma.docx", fakes_docx_mime(), caption="rus tiliga tarjima qil")
            create = env.api.requests[1]
            assert create.get("stream") is not True and create["model"] == "claude-opus-5-5"
            assert "tarjimon" in create["system"] and json.loads(create["messages"][0]["content"]) == ["Shartnoma", "Tomonlar kelishib oldilar."]
            name, data, _ = env.session.documents[-1]
            assert name == "shartnoma_tarjima.docx"
            texts = [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]
            assert texts == ["Договор", "Стороны договорились."]
            assert "2 ta abzats" in env.api.requests[2]["messages"][-1]["content"][0]["content"]
            final = env.last_edit()
            assert "Tarjima tayyor" in final.text and "📎 1" in final.text
            rows = await env.db.spend(OWNER, env.ctx.now().replace(hour=0, minute=0, second=0, microsecond=0))
            assert rows[1] == 2  # suhbat navbati (ikki oqim bitta yozuv) + tarjima so'rovi
        finally:
            await env.close()
    run(go())


def fakes_docx_mime() -> str:
    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_compare_documents_tool(tmp_path):
    async def go():
        env = await Env().start(tmp_path, [
            fakes.tool_answer("compare_documents", {}),
            fakes.simple_answer("Farq: ikkinchi faylda narx o'zgargan."),
        ])
        try:
            await env.send_document("t1", b"Narx: 100\nMuddat: 10 kun\n", "eski.txt", "text/plain")
            await env.send_document("t2", b"Narx: 120\nMuddat: 10 kun\n", "yangi.txt", "text/plain")
            await env.send_text("ikkalasini solishtir")
            result = env.api.requests[1]["messages"][-1]["content"][0]["content"]
            assert "-Narx: 100" in result and "+Narx: 120" in result and "Olib tashlangan qatorlar: 1" in result
            assert env.session.documents == []
            assert "narx o'zgargan" in env.last_edit().text
        finally:
            await env.close()
    run(go())


def test_tool_round_limit(tmp_path):
    async def go():
        answers = [
            fakes.tool_answer("create_document", {"format": "txt", "filename": f"f{i}", "content": "x"}, tool_id=f"toolu_{i}")
            for i in range(MAX_TOOL_ROUNDS + 1)
        ]
        env = await Env().start(tmp_path, answers)
        try:
            await env.send_text("ko'p fayl")
            assert len(env.api.requests) == MAX_TOOL_ROUNDS + 1
            assert len(env.session.documents) == MAX_TOOL_ROUNDS
            final = env.last_edit()
            assert "chegarasiga yetildi" in final.text and f"📎 {MAX_TOOL_ROUNDS}" in final.text
        finally:
            await env.close()
    run(go())


def test_quick_task_has_no_tools(tmp_path):
    async def go():
        env = await Env().start(tmp_path, [fakes.simple_answer("Uzun javob matni."), fakes.simple_answer("Qisqa.", model="claude-haiku-5-5")])
        try:
            await env.send_text("savol")
            await env.press("ans:short:2")
            assert "tools" not in env.api.requests[1]
        finally:
            await env.close()
    run(go())


# ------------------------------------------------------------------ birlik testlari


def test_raw_store_limits_and_ttl():
    store = RawStore(ttl=100, max_files=3, max_bytes=10)
    store.add(1, [RawFile("a", "text", b"1234"), RawFile("b", "text", b"1234")])
    store.add(1, [RawFile("c", "text", b"1234")])  # 12 bayt > 10 → eng eskisi chiqadi
    assert [f.name for f in store.recent(1)] == ["b", "c"]
    store.add(1, [RawFile("d", "text", b"1"), RawFile("e", "text", b"1")])
    assert [f.name for f in store.recent(1)] == ["c", "d", "e"]
    store._items[1][0].stamp = time.monotonic() - 1000
    assert [f.name for f in store.recent(1)] == ["d", "e"]
    assert store.recent(2) == []


def test_pick_by_name_and_kind():
    items = [RawFile("Xat.PDF", "pdf", b"%PDF"), RawFile("rasm-1.jpg", "image", b""), RawFile("jadval.xlsx", "xlsx", b"")]
    assert [f.name for f in tools.pick(items, ["xat"], ("pdf",))] == ["Xat.PDF"]
    assert [f.name for f in tools.pick(items, None, ("image",))] == ["rasm-1.jpg"]
    with pytest.raises(tools.ToolError, match="topilmadi"):
        tools.pick(items, ["yoq.pdf"], ("pdf",))
    with pytest.raises(tools.ToolError, match="mos emas"):
        tools.pick(items, ["jadval.xlsx"], ("pdf",))
    with pytest.raises(tools.ToolError, match="kamida 2"):
        tools.pick(items, None, ("pdf",), at_least=2)


def test_parse_pages():
    assert pdfs.parse_pages("1-3,5", 10) == [0, 1, 2, 4]
    assert pdfs.parse_pages(" 2 ", 2) == [1]
    assert pdfs.parse_pages(None, 3) == [0, 1, 2]
    with pytest.raises(pdfs.PdfError):
        pdfs.parse_pages("4-9", 3)
    with pytest.raises(pdfs.PdfError):
        pdfs.parse_pages("abc", 3)


def test_pdf_operations_roundtrip():
    three = files.build_pdf("\n\n".join(["Sahifa " + "x" * 3000] * 3), "Uch sahifa")
    assert pdfs.page_count(three) >= 3
    extracted, chosen = pdfs.extract(three, "1,3")
    assert pdfs.page_count(extracted) == 2 and chosen == [0, 2]
    parts = pdfs.split(extracted)
    assert len(parts) == 2 and all(pdfs.page_count(p) == 1 for p in parts)
    rotated = pdfs.rotate(extracted, "1", 90)
    assert PdfReader(io.BytesIO(rotated)).pages[0].rotation == 90
    merged = pdfs.merge([extracted, parts[0]])
    assert pdfs.page_count(merged) == 3
    assert "Matn qatlami: bor" in pdfs.info(merged, "m.pdf") and "[1-sahifa]" in pdfs.text(merged)
    with pytest.raises(pdfs.PdfError, match="PDF emas"):
        pdfs.page_count(b"hello")

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (300, 200), (10, 20, 30)).save(buffer, "JPEG")
    scan = pdfs.images_to_pdf([buffer.getvalue(), buffer.getvalue()])
    assert pdfs.page_count(scan) == 2


def test_docx_translation_keeps_structure():
    document = docx.Document()
    document.add_heading("Sarlavha", 1)
    paragraph = document.add_paragraph()
    bold = paragraph.add_run("Tomonlar: ")
    bold.bold = True
    paragraph.add_run("Alfa va Beta")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Nomi", "120 000"
    document.add_paragraph("2026-01-01")  # harf yo'q — tarjima qilinmaydi
    buffer = io.BytesIO()
    document.save(buffer)

    seen: list[list[str]] = []

    async def translate(batch):
        seen.append(batch)
        return [f"[ru] {t}" for t in batch]

    async def go():
        return await doctrans.translate_docx(buffer.getvalue(), translate)

    data, count = run(go())
    assert count == 3 and seen == [["Sarlavha", "Tomonlar: Alfa va Beta", "Nomi"]]
    out = docx.Document(io.BytesIO(data))
    assert [p.text for p in out.paragraphs] == ["[ru] Sarlavha", "[ru] Tomonlar: Alfa va Beta", "2026-01-01"]
    assert out.paragraphs[0].style.name.startswith("Heading")
    assert out.paragraphs[1].runs[0].bold is True
    assert out.tables[0].cell(0, 0).text == "[ru] Nomi" and out.tables[0].cell(0, 1).text == "120 000"
    assert doctrans.batches(["a" * 5000, "b" * 2000, "c"]) == [["a" * 5000], ["b" * 2000, "c"]]


def test_markdown_fallbacks_and_builders():
    sheets = files.tables_from_markdown("## Sotuv\n\n| Oy | Summa |\n|---|---|\n| Yan | 100 |\n\nmatn\n\n| a |\n|---|\n| 1 |")
    assert [s["name"] for s in sheets] == ["Sotuv", "Varaq2"]
    assert sheets[0]["columns"] == ["Oy", "Summa"] and sheets[0]["rows"] == [["Yan", "100"]]
    workbook = load_workbook(io.BytesIO(files.build_xlsx(sheets)))
    assert workbook["Sotuv"]["B2"].value == 100 and workbook["Sotuv"].freeze_panes == "A2"
    formula = load_workbook(io.BytesIO(files.build_xlsx([{"rows": [[1], ["=SUM(A1:A1)"]]}])))
    assert formula.active["A2"].value == "=SUM(A1:A1)" and formula.active["A2"].data_type == "f"

    slides = files.slides_from_markdown("# Kirish\n- bir\n- ikki\n# Xulosa\n1. uch")
    assert slides == [{"title": "Kirish", "bullets": ["bir", "ikki"]}, {"title": "Xulosa", "bullets": ["uch"]}]
    deck = Presentation(io.BytesIO(files.build_pptx(slides, "Loyiha")))
    assert len(deck.slides) == 3  # titul + 2
    assert files.safe_name("Hisobot: 2026/09?", "docx") == "Hisobot_ 2026_09.docx"
    with pytest.raises(files.FileError):
        files.build_chart({"type": "bar", "labels": ["a"], "series": [{"values": [1, 2]}]})
    assert files.build_chart({"type": "line", "labels": ["a", "b"], "series": [{"values": [1, 2]}]})[:4] == b"\x89PNG"
    with zipfile.ZipFile(io.BytesIO(files.build_docx("# A\n\nB", "T"))) as archive:
        assert "word/document.xml" in archive.namelist()
