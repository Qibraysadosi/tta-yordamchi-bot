"""Formatlash, narx hisobi, fayllar va sozlamalar uchun birlik sinovlari."""
import base64
import io
from html.parser import HTMLParser
from types import SimpleNamespace

import pytest
from PIL import Image

from bot import media, models, tgformat
from bot.config import ConfigError, load_settings

ALLOWED = {"b", "i", "s", "u", "code", "pre", "a", "blockquote", "tg-spoiler"}


class TagChecker(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=False)
        self.stack, self.errors = [], []

    def handle_starttag(self, tag, attrs):
        if tag not in ALLOWED:
            self.errors.append(f"ruxsatsiz teg: {tag}")
        self.stack.append(tag)

    def handle_endtag(self, tag):
        if not self.stack or self.stack[-1] != tag:
            self.errors.append(f"noto'g'ri yopilish: {tag} (stack={self.stack})")
        else:
            self.stack.pop()


def assert_valid_html(h: str):
    checker = TagChecker()
    checker.feed(h)
    checker.close()
    assert not checker.errors, checker.errors
    assert not checker.stack, f"yopilmagan: {checker.stack}"
    assert len(h) <= tgformat.TG_LIMIT


def test_inline_formatting():
    md = "**Salom** _dunyo_ va *kursiv*, `a < b && c`, [havola](https://kun.uz/?a=1&b=2) ~~eski~~"
    (h, p), = tgformat.format_answer(md, [], [], None)
    assert_valid_html(h)
    assert "<b>Salom</b>" in h and "<i>dunyo</i>" in h and "<i>kursiv</i>" in h
    assert "<code>a &lt; b &amp;&amp; c</code>" in h
    assert '<a href="https://kun.uz/?a=1&amp;b=2">havola</a>' in h
    assert "<s>eski</s>" in h
    assert p.startswith("**Salom**")


def test_no_false_italics():
    md = "snake_case_name va 2 * 3 * 4 = 24, fayl_nomi.txt, a*b*c"
    (h, _), = tgformat.format_answer(md, [], [], None)
    assert "<i>" not in h, h
    assert_valid_html(h)


def test_headers_lists_quotes_hr():
    md = "# Sarlavha **qalin**\n\n- birinchi\n  - ichki\n* ikkinchi\n\n> iqtibos **muhim**\n> davomi\n\n---\n1. raqamli"
    (h, _), = tgformat.format_answer(md, [], [], None)
    assert_valid_html(h)
    assert "<b>Sarlavha qalin</b>" in h
    assert "• birinchi" in h and "   • ichki" in h and "• ikkinchi" in h
    assert "<blockquote>iqtibos <b>muhim</b>\ndavomi</blockquote>" in h
    assert "──────────" in h and "1. raqamli" in h


def test_code_block_kept_and_escaped():
    md = "Kod:\n```python\nif a < b:\n    print('<html>')\n```\nTamom."
    (h, _), = tgformat.format_answer(md, [], [], None)
    assert_valid_html(h)
    assert '<pre><code class="language-python">if a &lt; b:\n    print(\'&lt;html&gt;\')</code></pre>' in h


def test_unclosed_fence_and_inline_triple():
    md = "```js\nconsole.log(1)"
    (h, _), = tgformat.format_answer(md, [], [], None)
    assert_valid_html(h)
    assert "console.log(1)" in h
    (h2, _), = tgformat.format_answer("Bu ```x``` inline", [], [], None)
    assert_valid_html(h2)


def test_tables():
    narrow = "| A | B |\n|---|---|\n| 1 | 2 |"
    (h, _), = tgformat.format_answer(narrow, [], [], None)
    assert_valid_html(h)
    assert h.startswith("<pre>") and "A │ B" in h
    wide = "| Model | Narx | Tavsif |\n|---|---|---|\n| **Opus 5.5** | $4 / $20 | kuchli va o'rtacha tezlikdagi model |"
    (h2, _), = tgformat.format_answer(wide, [], [], None)
    assert_valid_html(h2)
    assert "▪️ <b>Opus 5.5</b> — Narx: $4 / $20; Tavsif:" in h2


def test_long_answer_split_and_code_not_broken():
    para = "Bu juda uzun gap. " * 40
    md = "\n\n".join([para] * 12) + "\n```python\n" + "\n".join(f"print({i})" for i in range(300)) + "\n```"
    chunks = tgformat.format_answer(md, [("Manba", "https://a.uz")], ["eslatma"], "Opus 5.5 · $0.01")
    assert len(chunks) >= 3
    for h, p in chunks:
        assert h is not None
        assert_valid_html(h)
        assert len(p) <= tgformat.TG_LIMIT
    joined = "".join(h for h, _ in chunks)
    assert joined.count("<pre>") == joined.count("</pre>")
    assert "print(299)" in joined
    assert "Manbalar" in chunks[-1][0] and "<i>Opus 5.5 · $0.01</i>" in chunks[-1][0]


def test_sources_and_empty():
    chunks = tgformat.format_answer("", [], ["⏹ Javob to'xtatildi."], None)
    assert chunks == [("⏹ Javob to&#x27;xtatildi.", "⏹ Javob to'xtatildi.")] or chunks[0][1] == "⏹ Javob to'xtatildi."
    chunks = tgformat.format_answer("", [], [], None)
    assert chunks[0][1] == "(bo'sh javob)"
    (h, _), = tgformat.format_answer("Javob", [("Sarlavha <x>", "https://ex.com/a?b=1&c=2")], [], None)
    assert_valid_html(h)
    assert "<blockquote expandable>" in h and "Sarlavha &lt;x&gt;" in h


def test_preview():
    assert tgformat.preview("abc") == "abc"
    long = "x" * 5000
    assert tgformat.preview(long).startswith("…") and len(tgformat.preview(long)) == tgformat.PREVIEW_LIMIT + 1


def test_cost():
    usage = SimpleNamespace(
        input_tokens=1000, output_tokens=2000, cache_creation_input_tokens=10000, cache_read_input_tokens=50000,
        cache_creation=SimpleNamespace(ephemeral_5m_input_tokens=10000, ephemeral_1h_input_tokens=0),
        output_tokens_details=SimpleNamespace(thinking_tokens=1500),
        server_tool_use=SimpleNamespace(web_search_requests=2, web_fetch_requests=0),
    )
    n = models.usage_numbers(usage)
    assert n["thinking"] == 1500 and n["searches"] == 2 and n["cache_write_5m"] == 10000
    # Opus 5.5: 1000*4 + 2000*20 + 10000*5 + 50000*0.2 = 4000+40000+50000+10000 = 104000 /1e6 = 0.104 (+0.02 qidiruv)
    assert abs(models.cost_usd("claude-opus-5-5", n) - 0.124) < 1e-9
    # Fable 5.1: 1000*10 + 2000*50 + 10000*12.5 + 50000*0.25 = 10000+100000+125000+12500 = 0.2475 +0.02
    assert abs(models.cost_usd("claude-fable-5-1", n) - 0.2675) < 1e-9
    # Haiku 5.5, uzun prompt (>100k) narxi
    big = {"input": 150_000, "output": 1000}
    assert abs(models.cost_usd("claude-haiku-5-5", big) - (150_000 * 0.5 + 1000 * 2.5) / 1e6) < 1e-12
    small = {"input": 50_000, "output": 1000}
    assert abs(models.cost_usd("claude-haiku-5-5", small) - (50_000 * 0.1 + 1000 * 0.5) / 1e6) < 1e-12
    # 1 soatlik kesh: breakdown bo'lmasa sozlamaga qarab
    n1h = models.usage_numbers({"cache_creation_input_tokens": 1000}, "1h")
    assert n1h["cache_write_1h"] == 1000 and n1h["cache_write_5m"] == 0


def _png(size=(3000, 2000), mode="RGB"):
    buf = io.BytesIO()
    Image.new(mode, size, (200, 10, 10) if mode == "RGB" else (200, 10, 10, 120)).save(buf, "PNG")
    return buf.getvalue()


def test_media_images_and_files():
    block = media.image_block(_png(), 1568)
    assert block["source"]["media_type"] == "image/jpeg"
    img = Image.open(io.BytesIO(base64.b64decode(block["source"]["data"])))
    assert max(img.size) == 1568
    alpha = media.image_block(_png((400, 300), "RGBA"), 1568)
    assert alpha["source"]["media_type"] == "image/png"
    with pytest.raises(media.MediaError):
        media.image_block(b"not an image", 1568)

    assert media.classify("hisobot.PDF", "") == "pdf"
    assert media.classify("x.docx", "") == "docx"
    assert media.classify("x.xlsx", "application/octet-stream") == "xlsx"
    assert media.classify("main.py", "") == "text"
    assert media.classify("rasm.heic", "image/heic") is None
    assert media.classify("photo.jpg", "") == "image"

    pdf = media.pdf_block(b"%PDF-1.7\n...", "a.pdf")
    assert pdf["source"]["media_type"] == "application/pdf" and pdf["title"] == "a.pdf"
    with pytest.raises(media.MediaError):
        media.pdf_block(b"hello", "a.pdf")

    assert media.decode_text("Салом".encode("cp1251")) == "Салом"
    assert media.decode_text("\ufeffSalom".encode("utf-8")) == "Salom"
    tb = media.text_block("abc", "t.txt")
    assert tb["source"] == {"type": "text", "media_type": "text/plain", "data": "abc"}
    with pytest.raises(media.MediaError):
        media.text_block("   ", "t.txt")


def test_docx_xlsx():
    import docx
    import openpyxl

    d = docx.Document()
    d.add_paragraph("Shartnoma matni")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Ism", "Summa"
    t.cell(1, 0).text, t.cell(1, 1).text = "Ali", "100"
    t.cell(1, 0).merge(t.cell(1, 1))  # birlashtirilgan katak
    buf = io.BytesIO()
    d.save(buf)
    out = media.docx_text(buf.getvalue())
    assert "Shartnoma matni" in out and "Ism | Summa" in out
    assert out.count("Ali") == 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lidlar"
    ws.append(["Ism", "Telefon", None])
    ws.append(["Vali", "+998901234567", None])
    ws.append([None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    out = media.xlsx_text(buf.getvalue())
    assert out == "### Varaq: Lidlar\nIsm\tTelefon\nVali\t+998901234567"


def test_config(monkeypatch):
    for k in ("TELEGRAM_BOT_TOKEN", "ANTHROPIC_API_KEY", "OWNER_IDS"):
        monkeypatch.delenv(k, raising=False)
    with pytest.raises(ConfigError):
        load_settings()
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", "1:abc")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-x")
    monkeypatch.setenv("OWNER_IDS", "5591487950, 123")
    s = load_settings()
    assert s.owner_ids == frozenset({5591487950, 123}) and s.default_model == "claude-opus-5-5"
    assert s.fallback_model == "claude-opus-4-8" and s.cache_ttl == "5m" and s.max_tokens == 32000
    monkeypatch.setenv("FALLBACK_MODEL", "off")
    assert load_settings().fallback_model is None
    monkeypatch.setenv("DEFAULT_MODEL", "gpt-5")
    with pytest.raises(ConfigError):
        load_settings()
